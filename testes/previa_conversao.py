#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Prévia visual da conversão PDF -> PowerPoint.

Gera um PNG com o primeiro slide (e os seguintes) exatamente como o
PowerPoint/Holyrics vai mostrar.

  python testes/previa_conversao.py                 # usa PDFs de exemplo
  python testes/previa_conversao.py meu-arquivo.pdf # usa o seu PDF
  python testes/previa_conversao.py meu.pdf saida.png 3   # + destino e nº de slides
"""

from __future__ import annotations

import io
import json
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

import bot  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

LARG, ALT = 1280, 720  # 96 dpi do slide 16:9


def pdf_exemplo_letra(i: int, n: int) -> Image.Image:
    img = Image.new("RGB", (1920, 1080), (18, 22, 31))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, 1920, 8], fill=(214, 168, 84))
    d.text((960, 380), f"ESTROFE {i} DE {n}", fill=(255, 255, 255), anchor="mm")
    d.text((960, 460), "Tua graca me sustenta, Senhor", fill=(214, 168, 84), anchor="mm")
    d.rectangle([200, 900, 1720, 906], fill=(60, 70, 90))
    d.text((960, 1000), "Culto de Domingo - 19h", fill=(140, 150, 165), anchor="mm")
    return img


def pdf_exemplo_a4() -> Image.Image:
    img = Image.new("RGB", (1240, 1754), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((80, 80), "ESCALA DE MIDIA - documento A4 retrato", fill=(20, 20, 20))
    for k in range(14):
        d.text((80, 220 + k * 70), f"{k + 9}h - Equipe {k + 1} - som e projecao", fill=(60, 60, 60))
    return img


def render_slide(pptx: Path, indice: int) -> Image.Image:
    """Monta a imagem do slide como o PowerPoint desenha (fundo branco + fotos)."""
    prs = Presentation(str(pptx))
    tela = Image.new("RGB", (LARG, ALT), (255, 255, 255))
    for sh in list(prs.slides)[indice].shapes:
        if sh.shape_type != 13:  # 13 = picture
            continue
        foto = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
        x, y = Emu(sh.left).inches * 96, Emu(sh.top).inches * 96
        w, h = Emu(sh.width).inches * 96, Emu(sh.height).inches * 96
        foto = foto.resize((max(1, round(w)), max(1, round(h))), Image.LANCZOS)
        tela.paste(foto, (round(x), round(y)))
    return tela


def main() -> int:
    args = [a for a in sys.argv[1:]]
    pdf_usuario = Path(args[0]) if args and not args[0].isdigit() else None
    saida = Path(args[1]) if len(args) > 1 else RAIZ / "testes" / "previa-conversao.png"
    quantos = int(args[2]) if len(args) > 2 else 2

    tmp = Path(tempfile.mkdtemp(prefix="previa-"))
    cfg_path = tmp / "config.json"
    cfg_path.write_text(
        json.dumps({"telegram": {"token": "x"}, "pastas": {"apresentacao": str(tmp / "saida")}}),
        encoding="utf-8",
    )
    cfg = bot.Config(cfg_path)
    pasta = cfg.pastas["apresentacao"]
    pasta.mkdir(parents=True, exist_ok=True)

    if pdf_usuario:
        arquivos = [(pdf_usuario, pasta / (pdf_usuario.stem + ".pptx"), pdf_usuario.name)]
    else:
        p1 = tmp / "letras-16x9.pdf"
        paginas = [pdf_exemplo_letra(i, 3) for i in (1, 2, 3)]
        paginas[0].save(p1, "PDF", save_all=True, append_images=paginas[1:], resolution=150)
        p2 = tmp / "escala-A4.pdf"
        pdf_exemplo_a4().save(p2, "PDF", resolution=150)
        arquivos = [(p1, pasta / "letras-16x9.pptx", "exemplo 16:9"), (p2, pasta / "escala-A4.pptx", "exemplo A4 retrato")]

    quadros = []
    for pdf, destino, rotulo in arquivos:
        n = bot.pdf_para_pptx(pdf, destino, cfg)
        for i in range(min(quantos, n)):
            tela = render_slide(destino, i)
            faixa = Image.new("RGB", (LARG, 34), (24, 28, 36))
            ImageDraw.Draw(faixa).text(
                (12, 10),
                f"{destino.name}  -  slide {i + 1}/{n}  -  {destino.stat().st_size // 1024} KB  ({rotulo})",
                fill=(220, 220, 220),
            )
            tela.paste(faixa, (0, 0))
            quadros.append(tela)

    final = Image.new("RGB", (LARG, ALT * len(quadros) + 10 * (len(quadros) - 1)), (190, 190, 190))
    for i, q in enumerate(quadros):
        final.paste(q, (0, i * (ALT + 10)))
    saida.parent.mkdir(parents=True, exist_ok=True)
    final.save(saida)
    print(f"Prévia salva em: {saida}  ({len(quadros)} slide(s) mostrado(s))")
    return 0


if __name__ == "__main__":
    sys.exit(main())
