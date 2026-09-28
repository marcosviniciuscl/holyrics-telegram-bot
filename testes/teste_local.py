#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Teste local do bot, sem precisar do Telegram.
Valida:
  1) a conversão PDF -> PowerPoint (.pptx)
  2) o download em streaming com atualização de progresso (sem spam)

Uso:  python testes/teste_local.py
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import sys
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

import bot  # noqa: E402

TAMANHO_TESTE = 5 * 1024 * 1024  # 5 MB
CONTEUDO = bytes(range(256)) * (TAMANHO_TESTE // 256)


def _cfg_minimo(tmp: Path, nome: str, ffmpeg: str = "") -> Path:
    """Config enxuto para os testes."""
    caminho = tmp / nome
    caminho.write_text(json.dumps({
        "telegram": {"token": "x", "usuarios_autorizados": [7]},
        "pastas": {"video": str(tmp / "Videos"), "audio": str(tmp / "Audios")},
        "youtube": {"ativo": True, "ffmpeg": ffmpeg},
    }), encoding="utf-8")
    return caminho


def criar_pdf(caminho: Path, paginas: int = 5) -> Path:
    from PIL import Image, ImageDraw

    imagens = []
    for i in range(1, paginas + 1):
        img = Image.new("RGB", (1600, 900), "#101820")
        d = ImageDraw.Draw(img)
        d.rectangle([60, 60, 1540, 840], outline="#f2aa4c", width=6)
        d.text((120, 400), f"SLIDE {i} DE {paginas}", fill="white")
        imagens.append(img)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    imagens[0].save(caminho, "PDF", save_all=True, append_images=imagens[1:], resolution=150)
    return caminho


def teste_conversao(tmp: Path) -> bool:
    print("=" * 62)
    print("1) CONVERSÃO PDF -> POWERPOINT")
    print("=" * 62)

    config = {
        "telegram": {"token": "teste", "usuarios_autorizados": [1]},
        "pastas": {
            "apresentacao": str(tmp / "Apresentacoes"),
            "audio": str(tmp / "Audios"),
        },
        "pdf": {"converter_para_pptx": True, "dpi": 200, "largura_max_px": 2400},
    }
    caminho_cfg = tmp / "config.json"
    caminho_cfg.write_text(json.dumps(config), encoding="utf-8")
    cfg = bot.Config(caminho_cfg)

    pdf = criar_pdf(tmp / "liturgia.pdf", paginas=5)
    print(f"PDF de teste: {pdf.name} — {bot.fmt_bytes(pdf.stat().st_size)}")

    destino_pasta = cfg.pastas["apresentacao"]
    destino_pasta.mkdir(parents=True, exist_ok=True)
    pptx = destino_pasta / "liturgia.pptx"

    import time

    t0 = time.time()
    paginas = bot.pdf_para_pptx(pdf, pptx, cfg)
    dur = time.time() - t0

    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(str(pptx))
    slides = list(prs.slides)
    fotos = [len([s for s in sl.shapes if s.shape_type == 13]) for sl in slides]

    print(f"Páginas convertidas: {paginas} em {dur:.2f}s")
    print(f"Slides no .pptx: {len(slides)}")
    print(f"Imagens por slide: {fotos}")
    print(
        f"Tamanho do slide: {Emu(prs.slide_width).inches:.3f} x "
        f"{Emu(prs.slide_height).inches:.3f} polegadas (16:9 = 13.333 x 7.500)"
    )
    print(f"Tamanho do .pptx: {bot.fmt_bytes(pptx.stat().st_size)}")

    ok = (
        paginas == 5
        and len(slides) == 5
        and all(f == 1 for f in fotos)
        and abs(Emu(prs.slide_width).inches - 13.333) < 0.01
    )
    print("RESULTADO:", "OK ✅" if ok else "FALHOU ❌")
    return ok


class Servidor(threading.Thread):
    """Servidor HTTP local que serve bytes fixos, para testar o download."""

    def __init__(self, conteudo: bytes, content_type: str = "application/octet-stream"):
        super().__init__(daemon=True)
        corpo = conteudo

        class Handler(BaseHTTPRequestHandler):
            def _cabecalhos(self):
                self.send_response(200)
                self.send_header("Content-Type", content_type)
                self.send_header("Content-Length", str(len(corpo)))
                self.send_header("Accept-Ranges", "none")
                self.end_headers()

            def do_GET(self):  # noqa: N802
                self._cabecalhos()
                # envia em blocos, como o Telegram faria
                try:
                    for i in range(0, len(corpo), 64 * 1024):
                        self.wfile.write(corpo[i : i + 64 * 1024])
                except (ConnectionResetError, BrokenPipeError):
                    pass  # o cliente fechou antes (normal em testes)

            def do_HEAD(self):  # noqa: N802
                self._cabecalhos()

            def log_message(self, *a):
                pass

            def handle_one_request(self):
                try:
                    super().handle_one_request()
                except (ConnectionResetError, BrokenPipeError):
                    self.close_connection = True

        self.httpd = HTTPServer(("127.0.0.1", 0), Handler)
        self.porta = self.httpd.server_port

    def run(self):
        self.httpd.serve_forever()

    def parar(self):
        self.httpd.shutdown()


class MensagemFalsa:
    """Imita a mensagem do Telegram, guardando cada edição de texto."""

    def __init__(self):
        self.edicoes: list[str] = []

    async def edit_text(self, texto, **kwargs):
        self.edicoes.append(texto)


def teste_download(tmp: Path) -> bool:
    import asyncio

    print()
    print("=" * 62)
    print("2) DOWNLOAD COM PROGRESSO (mensagem sendo editada, sem spam)")
    print("=" * 62)

    servidor = Servidor(CONTEUDO)
    servidor.start()
    try:
        msg = MensagemFalsa()
        progresso = bot.Progresso(
            msg, "<b>video.mp4</b>", TAMANHO_TESTE, intervalo=0.0  # 0s só para o teste
        )
        destino = tmp / "video.mp4"
        feito = asyncio.run(
            bot.baixar(f"http://127.0.0.1:{servidor.porta}/video.mp4", destino, progresso)
        )
        sha_origem = hashlib.sha256(CONTEUDO).hexdigest()
        sha_destino = hashlib.sha256(destino.read_bytes()).hexdigest()
    finally:
        servidor.parar()

    print(f"Bytes baixados: {feito} de {TAMANHO_TESTE}")
    print(f"SHA256 confere: {sha_origem == sha_destino}")
    print(f"Edições na mensagem: {len(msg.edicoes)} (progresso vai até 99%, sem 100% duplicado)")
    for e in msg.edicoes:
        print("   •", e.replace("\n", " | "))
    print(f"Tamanho final: {bot.fmt_bytes(destino.stat().st_size)}")

    pcts = []
    for e in msg.edicoes:
        if "%" in e:
            pcts.append(int(e.split("…")[1].split("%")[0].strip()))
    crescente = all(b >= a for a, b in zip(pcts, pcts[1:]))
    ok = (
        sha_origem == sha_destino
        and feito == TAMANHO_TESTE
        and crescente
        and len(pcts) <= 30
        and not destino.with_name(destino.name + ".parcial").exists()
    )
    print("RESULTADO:", "OK ✅" if ok else "FALHOU ❌")
    return ok


def teste_regras(tmp: Path) -> bool:
    """Regras de categoria, nome seguro e nomes repetidos."""
    from types import SimpleNamespace

    print()
    print("=" * 62)
    print("3) REGRAS DE ORGANIZAÇÃO (categoria, nome, duplicados)")
    print("=" * 62)

    categorias = {
        "missa.mp3": "audio",
        "clipe.MKV": "video",
        "foto.JPEG": "imagem",
        "letra.PPTX": "apresentacao",
        "escala.docx": "documento",
        "liturgia.pdf": "outros",  # PDF é tratado separadamente
        "backup.zip": "outros",
    }
    linhas = []
    ok_cat = True
    for nome, esperado in categorias.items():
        obtido = bot.categoria_de(nome)
        ok_cat &= obtido == esperado
        linhas.append(f"  {nome:16} -> {obtido:13} {'✅' if obtido == esperado else '❌ esperado ' + esperado}")
    print("\n".join(linhas))

    print("Nome seguro:")
    for entrada, esperado in [
        ("../../etc/passwd", "passwd"),
        ("letra: ruim?.pptx", "letra_ ruim_.pptx"),
        ("", "arquivo"),
    ]:
        obtido = bot.nome_seguro(entrada)
        ok_cat &= obtido == esperado
        print(f"  {entrada!r:22} -> {obtido!r} {'✅' if obtido == esperado else '❌'}")

    pasta = tmp / "dup"
    pasta.mkdir(parents=True, exist_ok=True)
    (pasta / "letra.mp3").write_bytes(b"a")
    n2 = bot.nome_livre(pasta / "letra.mp3")
    (pasta / "letra (2).mp3").write_bytes(b"b")
    n3 = bot.nome_livre(pasta / "letra.mp3")
    print(f"Duplicados: letra.mp3 -> {n2.name} -> {n3.name}")
    ok_cat &= n2.name == "letra (2).mp3" and n3.name == "letra (3).mp3"

    # extrair_arquivo para cada tipo de mensagem do Telegram
    print("Detecção por tipo de mensagem:")
    dados = [
        (SimpleNamespace(document=SimpleNamespace(file_id="a", file_name="m.pdf", file_size=10, mime_type="application/pdf")), "m.pdf"),
        (SimpleNamespace(document=SimpleNamespace(file_id="b", file_name=None, file_size=10, mime_type="audio/mpeg")), ".mp3"),
        (SimpleNamespace(document=None, audio=SimpleNamespace(file_id="c", file_name="hino.wav", file_size=10, mime_type="audio/wav")), "hino.wav"),
        (SimpleNamespace(document=None, audio=None, video=None, voice=SimpleNamespace(file_id="d", file_size=10, mime_type="audio/ogg")), ".ogg"),
        (SimpleNamespace(document=None, audio=None, video=None, voice=None, video_note=SimpleNamespace(file_id="e", file_size=10)), ".mp4"),
        (SimpleNamespace(document=None, audio=None, video=None, voice=None, video_note=None, animation=None,
                         photo=[SimpleNamespace(file_id="f", file_size=10), SimpleNamespace(file_id="g", file_size=99)]), ".jpg"),
    ]
    for msg, esperado in dados:
        for campo in ("document", "audio", "video", "voice", "video_note", "animation", "photo", "sticker"):
            if not hasattr(msg, campo):
                setattr(msg, campo, None)
        r = bot.extrair_arquivo(msg)
        obtido = r[1] if r else None
        casou = bool(obtido) and (obtido.endswith(esperado) if esperado.startswith(".") else obtido == esperado)
        ok_cat &= casou
        print(f"  {esperado:10} -> {obtido} {'✅' if casou else '❌'}")

    # foto: deve escolher a maior resolução
    ok_cat &= bot.extrair_arquivo(dados[-1][0])[0] == "g"
    print("Foto escolhe a maior resolução:", bot.extrair_arquivo(dados[-1][0])[0] == "g")

    print("RESULTADO:", "OK ✅" if ok_cat else "FALHOU ❌")
    return ok_cat


class BotFalso:
    """Imita a API de arquivos do Telegram (aponta para o servidor local)."""

    def __init__(self, file_path: str, file_size: int):
        self.file_path = file_path
        self.file_size = file_size

    async def get_file(self, file_id):
        from types import SimpleNamespace

        return SimpleNamespace(file_path=self.file_path, file_size=self.file_size)

    async def send_chat_action(self, *a, **k):
        return True


class ContextoFalso:
    def __init__(self, cfg, bot_falso):
        self.bot = bot_falso
        self.bot_data = {"cfg": cfg}


class MensagemRecebidaFalsa:
    """Mensagem recebida do Telegram + a mensagem de status que o bot edita."""

    def __init__(self, documento, chat_id=1):
        self.document = documento
        self.text = None
        self.chat_id = chat_id
        self.chat = type("Chat", (), {"type": "private", "id": chat_id})()
        self.edicoes: list[str] = []

    async def reply_text(self, texto, **kwargs):
        self.edicoes.append(texto)
        return self

    async def edit_text(self, texto, **kwargs):
        self.edicoes.append(texto)

    async def edit_reply_markup(self, markup=None, **kwargs):
        self.markup = markup


def teste_ponta_a_ponta(tmp: Path) -> bool:
    """Fluxo completo: mensagem do Telegram -> download -> conversão -> mensagem final."""
    import asyncio
    from types import SimpleNamespace

    print()
    print("=" * 62)
    print("4) FLUXO COMPLETO (mensagem -> download -> PDF->PPTX -> resposta)")
    print("=" * 62)

    config = {
        "telegram": {"token": "teste", "usuarios_autorizados": [7]},
        "pastas": {
            "apresentacao": str(tmp / "Holyrics" / "Apresentacoes"),
            "audio": str(tmp / "Holyrics" / "Audios"),
            "outros": str(tmp / "Holyrics" / "Outros"),
        },
        "pdf": {"converter_para_pptx": True, "dpi": 150},
        "download": {"intervalo_progresso_segundos": 0},
    }
    cfg_path = tmp / "config-e2e.json"
    cfg_path.write_text(json.dumps(config), encoding="utf-8")
    cfg = bot.Config(cfg_path)

    # "arquivo do Telegram" servido pelo servidor local
    pdf = criar_pdf(tmp / "culto.pdf", paginas=4)
    conteudo = pdf.read_bytes()
    servidor = Servidor(conteudo)
    servidor.start()
    bot.BASE_API = f"http://127.0.0.1:{servidor.porta}"

    msg = MensagemRecebidaFalsa(
        SimpleNamespace(file_id="id1", file_name="culto.pdf", file_size=len(conteudo), mime_type="application/pdf")
    )
    update = SimpleNamespace(effective_message=msg, effective_user=SimpleNamespace(id=7))
    contexto = ContextoFalso(cfg, BotFalso("docs/culto.pdf", len(conteudo)))

    try:
        asyncio.run(bot.receber(update, contexto))
    finally:
        servidor.parar()
        bot.BASE_API = "https://api.telegram.org"

    print(f"Mensagens enviadas: {len(msg.edicoes)}")
    for e in msg.edicoes:
        print("   •", e.replace("\n", " | "))

    destino = Path(config["pastas"]["apresentacao"]) / "culto.pptx"
    final = msg.edicoes[-1]
    ok = (
        destino.is_file()
        and "culto.pptx" in final
        and "Convertido para PowerPoint (4 slides)" in final
        and str(destino) in final
        and not (destino.parent / "culto.pdf").exists()  # PDF original descartado
        and len(msg.edicoes) <= 8  # sem spam de mensagens
    )
    if destino.is_file():
        from pptx import Presentation

        print(f"Arquivo gerado: {destino} ({bot.fmt_bytes(destino.stat().st_size)}, "
              f"{len(Presentation(str(destino)).slides._sldIdLst)} slides)")

    # segundo arquivo, agora um áudio, para conferir o roteamento por categoria
    msg2 = MensagemRecebidaFalsa(
        SimpleNamespace(file_id="id2", file_name="../hino.mp3", file_size=len(conteudo), mime_type="audio/mpeg")
    )
    servidor2 = Servidor(conteudo)
    servidor2.start()
    bot.BASE_API = f"http://127.0.0.1:{servidor2.porta}"
    try:
        asyncio.run(
            bot.receber(
                SimpleNamespace(effective_message=msg2, effective_user=SimpleNamespace(id=7)),
                ContextoFalso(cfg, BotFalso("docs/hino.mp3", len(conteudo))),
            )
        )
    finally:
        servidor2.parar()
        bot.BASE_API = "https://api.telegram.org"

    destino_audio = Path(config["pastas"]["audio"]) / "hino.mp3"
    print("Áudio:", msg2.edicoes[-1].replace("\n", " | "))
    ok &= destino_audio.is_file() and "Audios" in msg2.edicoes[-1]

    # usuário não autorizado é bloqueado
    msg3 = MensagemRecebidaFalsa(
        SimpleNamespace(file_id="id3", file_name="x.mp3", file_size=10, mime_type="audio/mpeg")
    )
    asyncio.run(
        bot.receber(
            SimpleNamespace(effective_message=msg3, effective_user=SimpleNamespace(id=999)),
            ContextoFalso(cfg, BotFalso("docs/x.mp3", 10)),
        )
    )
    bloqueado = "Não autorizado" in (msg3.edicoes[0] if msg3.edicoes else "")
    print("Usuário não autorizado bloqueado:", bloqueado)
    ok &= bloqueado

    print("RESULTADO:", "OK ✅" if ok else "FALHOU ❌")
    return ok


def criar_mp4(caminho: Path, segundos: int = 12) -> Path | None:
    """Gera um MP4 pequeno com ffmpeg (usado como 'vídeo do YouTube' no teste)."""
    import shutil
    import subprocess

    if not shutil.which("ffmpeg"):
        return None
    subprocess.run(
        [
            "ffmpeg", "-y", "-loglevel", "error",
            "-f", "lavfi", "-i", "testsrc=size=640x360:rate=15",
            "-f", "lavfi", "-i", "sine=frequency=440",
            "-t", str(segundos), "-shortest",
            "-c:v", "libx264", "-preset", "ultrafast", "-pix_fmt", "yuv420p",
            "-c:a", "aac", "-b:a", "96k", "-movflags", "+faststart",
            str(caminho),
        ],
        check=True,
    )
    return caminho


class ConsultaFalsa:
    """Imita o clique no botão (callback query) do Telegram."""

    def __init__(self, data: str, user_id: int, mensagem):
        self.data = data
        self.from_user = type("U", (), {"id": user_id, "full_name": "Teste"})()
        self.message = mensagem
        self.alertas: list[str] = []
        self.respostas: list[str] = []

    async def answer(self, texto=None, show_alert=False):
        if texto:
            (self.alertas if show_alert else self.respostas).append(texto)


def teste_youtube(tmp: Path) -> bool:
    """Links, montagem das 2+2 opções e o download real (yt-dlp + ffmpeg)."""
    import asyncio
    from types import SimpleNamespace

    print()
    print("=" * 62)
    print("5) YOUTUBE (links, opções 2 vídeo + 2 áudio, download real)")
    print("=" * 62)

    ok = True

    # --- 5.1 links ---------------------------------------------------------
    casos = [
        ("https://www.youtube.com/watch?v=aqz-KE-bpKQ", "aqz-KE-bpKQ"),
        ("https://youtu.be/aqz-KE-bpKQ?si=xyz", "aqz-KE-bpKQ"),
        ("https://m.youtube.com/watch?v=aqz-KE-bpKQ&list=RD1&index=2", "aqz-KE-bpKQ"),
        ("https://music.youtube.com/watch?v=aqz-KE-bpKQ", "aqz-KE-bpKQ"),
        ("https://www.youtube.com/shorts/aqz-KE-bpKQ", "aqz-KE-bpKQ"),
        ("https://youtube.com/live/aqz-KE-bpKQ?feature=share", "aqz-KE-bpKQ"),
        ("olha esse vídeo https://youtu.be/aqz-KE-bpKQ valeu", "aqz-KE-bpKQ"),
    ]
    for texto, vid in casos:
        obtido = bot.achar_videos_youtube(texto)
        esperado = [f"https://www.youtube.com/watch?v={vid}"]
        casou = obtido == esperado
        ok &= casou
        print(f"  {'✅' if casou else '❌'} {texto[:58]:58} -> {obtido[0] if obtido else None}")
    for texto in ["https://vimeo.com/12345", "bom dia, tudo bem?", "https://youtube.com/@canal"]:
        casou = bot.achar_videos_youtube(texto) == []
        ok &= casou
        print(f"  {'✅' if casou else '❌'} ignora: {texto}")

    # --- 5.2 montagem das opções ------------------------------------------
    cfg_path = tmp / "config-yt.json"
    cfg_path.write_text(json.dumps({
        "telegram": {"token": "x", "usuarios_autorizados": [7]},
        "pastas": {
            "video": str(tmp / "Holyrics" / "Videos"),
            "audio": str(tmp / "Holyrics" / "Audios"),
        },
        "youtube": {"ativo": True, "altura_maxima": 1080, "mp3_bitrate_kbps": 192},
    }), encoding="utf-8")
    cfg = bot.Config(cfg_path)

    info = {
        "title": "Hino Teste — Ação de Graças",
        "uploader": "Canal Teste",
        "duration": 253,
        "formats": [
            {"format_id": "137", "ext": "mp4", "height": 1080, "vcodec": "avc1", "acodec": "none", "tbr": 4500, "filesize": 142_000_000},
            {"format_id": "136", "ext": "mp4", "height": 720, "vcodec": "avc1", "acodec": "none", "tbr": 2500, "filesize": 79_000_000},
            {"format_id": "135", "ext": "mp4", "height": 480, "vcodec": "avc1", "acodec": "none", "tbr": 1200, "filesize": 38_000_000},
            {"format_id": "18", "ext": "mp4", "height": 360, "vcodec": "avc1", "acodec": "mp4a", "tbr": 700, "filesize": 22_000_000},
            {"format_id": "140", "ext": "m4a", "abr": 128, "vcodec": "none", "acodec": "mp4a", "tbr": 128, "filesize": 4_100_000},
            {"format_id": "251", "ext": "webm", "abr": 160, "vcodec": "none", "acodec": "opus", "tbr": 160, "filesize": 5_100_000},
        ],
    }
    texto, linhas, selecoes = bot.montar_opcoes(info, cfg)
    print("\nMensagem de opções enviada no Telegram:")
    for l in texto.split("\n"):
        print("   |", l)
    print("Botões:")
    for linha in linhas:
        print("   |", "   ".join(etiqueta for etiqueta, _ in linha))

    ok &= len(linhas) == 2 and all(len(l) == 2 for l in linhas)
    ok &= selecoes[0]["rotulo"] == "1080p" and selecoes[1]["rotulo"] == "720p"
    ok &= selecoes[0]["tipo"] == "video" and selecoes[1]["tipo"] == "video"
    ok &= "height<=720" in selecoes[1]["formato"]
    # O Holyrics prefere MP3: as duas opções de áudio viram MP3.
    ok &= selecoes[2]["tipo"] == "audio" and selecoes[2]["mp3"] is True
    ok &= selecoes[2]["rotulo"] == "MP3 192 kbps"
    ok &= selecoes[3]["mp3"] is True and selecoes[3]["rotulo"] == "MP3 128 kbps"
    ok &= all(etq.strip() for linha in linhas for etq, _ in linha)
    ok &= sum(1 for linha in linhas for _ in linha) == 4  # exatamente 4 opções
    print("Estrutura (2 vídeo + 2 áudio, rótulos e seletores):", "OK ✅" if ok else "FALHOU ❌")

    # sem ffmpeg: não oferece MP3 e avisa
    original = bot.ffmpeg_disponivel
    bot.ffmpeg_disponivel = lambda _cfg: False
    texto_sem, linhas_sem, selecoes_sem = bot.montar_opcoes(info, cfg)
    bot.ffmpeg_disponivel = original
    sem_ffmpeg_ok = (
        "ffmpeg não encontrado" in texto_sem
        and all(not s.get("mp3") for s in selecoes_sem)
        and len(linhas_sem) == 2
    )
    ok &= sem_ffmpeg_ok
    print("Sem ffmpeg (avisa e não oferece MP3):", "OK ✅" if sem_ffmpeg_ok else "FALHOU ❌")

    # --- 5.3 download real (yt-dlp + ffmpeg) ------------------------------
    if criar_mp4(tmp / "video-teste.mp4") is None:
        print("ffmpeg ausente: pulando o teste de download real")
        print("RESULTADO:", "OK ✅" if ok else "FALHOU ❌")
        return ok

    conteudo = (tmp / "video-teste.mp4").read_bytes()
    print(f"\nVídeo de teste: {bot.fmt_bytes(len(conteudo))} (servido localmente, baixado pelo yt-dlp)")
    servidor = Servidor(conteudo, content_type="video/mp4")
    servidor.start()
    url = f"http://127.0.0.1:{servidor.porta}/video-teste.mp4"

    try:
        # 5.3.a pipeline direto: download + conversão para MP3
        estado: dict = {"fase": "iniciando", "feito": 0, "total": 0}
        pasta_audio = cfg.pastas["audio"]
        destino, _ = bot.baixar_youtube(
            estado, url, pasta_audio,
            {"formato": "bestaudio/best", "mp3": True, "rotulo": "MP3 192 kbps", "tipo": "audio"},
            cfg,
        )
        print("Progresso final do estado:", {k: v for k, v in estado.items() if k != "pp"})
        hooks_ok = bool(estado.get("feito")) and bool(estado.get("total")) and estado.get("fase") in ("salvando", "convertendo", "processando")
        ok &= hooks_ok
        print("Hooks de progresso do yt-dlp dispararam:", "OK ✅" if hooks_ok else "FALHOU ❌")
        print("Textos de progresso possíveis:")
        for fase, pp in (("baixando", None), ("processando", None), ("convertendo", "FFmpegExtractAudio"), ("salvando", None)):
            estado_prov = {"fase": fase, "pp": pp, "feito": 600000, "total": 1200000, "vel": 950000, "eta": 3}
            print("   |", bot.texto_progresso_yt(estado_prov, "Hino Teste", "MP3 192 kbps").replace("\n", " / "))
        ok &= destino.is_file() and destino.suffix == ".mp3"
        print(f"Arquivo convertido: {destino.name} ({bot.fmt_bytes(destino.stat().st_size)})")

        # 5.3.b fluxo dos botões: opções -> clique -> download -> mensagem final
        bot.info_youtube = lambda _url, _cfg: {**info, "title": "Video Local", "duration": 12}
        pendentes = {}
        for i, (rotulo, chave, mp3) in enumerate([("Melhor disponível", "video", False), ("MP3 192 kbps", "audio", True)]):
            token = f"tok{i}"
            pendentes[token] = {
                "url": url, "info": bot.info_youtube(url, cfg), "user": 7, "chat": 1,
                "selecoes": [{"formato": "bestaudio/best" if mp3 else "best", "mp3": mp3,
                              "rotulo": rotulo, "tipo": "audio" if mp3 else "video"}],
                "quando": 0,
            }
            mensagem = MensagemRecebidaFalsa(SimpleNamespace())
            consulta = ConsultaFalsa(f"yt|{token}|0", 7, mensagem)
            contexto = ContextoFalso(cfg, BotFalso("x", 1))
            contexto.bot_data["yt_pendentes"] = pendentes
            asyncio.run(bot.botao_youtube(SimpleNamespace(callback_query=consulta, effective_user=SimpleNamespace(id=7)), contexto))

            print(f"\n[{rotulo}] mensagens editadas: {len(mensagem.edicoes)}")
            for e in mensagem.edicoes:
                print("   |", e.replace("\n", " / "))
            esperado_dir = "Videos" if chave == "video" else "Audios"
            final = mensagem.edicoes[-1]
            casou = (
                "✅" in final and esperado_dir in final
                and len(mensagem.edicoes) >= 2
                and consulta.data.split("|")[1] not in contexto.bot_data["yt_pendentes"]  # token consumido
            )
            ok &= casou
            salvos = list((tmp / "Holyrics" / esperado_dir).glob("*"))
            print(f"   arquivos em {esperado_dir}: {[p.name for p in salvos]}")
            ok &= bool(salvos)

        # 5.3.c dono do link é protegido
        token = "tokdono"
        pendentes[token] = {
            "url": url, "info": info, "user": 999, "chat": 1,
            "selecoes": [{"formato": "best", "mp3": False, "rotulo": "720p", "tipo": "video"}], "quando": 0,
        }
        msg_x = MensagemRecebidaFalsa(SimpleNamespace())
        consulta = ConsultaFalsa(f"yt|{token}|0", 7, msg_x)
        contexto = ContextoFalso(cfg, BotFalso("x", 1))
        contexto.bot_data["yt_pendentes"] = pendentes
        asyncio.run(bot.botao_youtube(SimpleNamespace(callback_query=consulta, effective_user=SimpleNamespace(id=7)), contexto))
        bloqueado = bool(consulta.alertas) and "outro usuário" in consulta.alertas[0]
        ok &= bloqueado
        print("\nBotão de link de outro usuário bloqueado:", bloqueado)
    finally:
        servidor.parar()

    print("RESULTADO:", "OK ✅" if ok else "FALHOU ❌")
    return ok


def teste_ffmpeg(tmp: Path) -> bool:
    """Ordem de procura do ffmpeg e detecção de plataforma do instalador."""
    import platform as _platform

    import baixar_ffmpeg

    print()
    print("=" * 62)
    print("6) FFMPEG (ordem de procura + instalador portátil)")
    print("=" * 62)

    ok = True
    bin_falso = tmp / "bot-com-bin"
    (bin_falso / "bin").mkdir(parents=True, exist_ok=True)
    exe = bin_falso / "bin" / ("ffmpeg.exe" if os.name == "nt" else "ffmpeg")
    exe.write_text("#!/bin/sh\necho 'ffmpeg version 9.9-fake'\n", encoding="utf-8")
    exe.chmod(0o755)

    cfg_sem = bot.Config(_cfg_minimo(tmp, "cfg-sem-ffmpeg.json"))
    cfg_com = bot.Config(_cfg_minimo(tmp, "cfg-com-ffmpeg.json", ffmpeg="/usr/bin/ffmpeg"))
    cfg_errado = bot.Config(_cfg_minimo(tmp, "cfg-ffmpeg-errado.json", ffmpeg="/nao/existe/ffmpeg"))

    raiz_original = bot.RAIZ

    # 1) bin/ do bot tem prioridade sobre o PATH
    bot.RAIZ = bin_falso
    r1 = bot.caminho_ffmpeg(cfg_sem)
    print(f"  bin/ portátil encontrado: {r1}")
    ok &= r1 == str(exe) and bot.ffmpeg_disponivel(cfg_sem)

    # 2) config.json tem prioridade sobre tudo
    r2 = bot.caminho_ffmpeg(cfg_com)
    print(f"  youtube.ffmpeg respeitado: {r2}")
    ok &= r2 == "/usr/bin/ffmpeg"

    # 3) sem bin/ e sem config -> PATH do sistema
    bot.RAIZ = tmp / "bot-sem-bin"
    (bot.RAIZ).mkdir(parents=True, exist_ok=True)
    r3 = bot.caminho_ffmpeg(cfg_sem)
    print(f"  cai para o PATH: {r3}")
    ok &= bool(r3) and r3 == shutil.which("ffmpeg")

    # 4) caminho inválido no config não quebra: avisa e usa o próximo
    r4 = bot.caminho_ffmpeg(cfg_errado)
    print(f"  youtube.ffmpeg inválido -> usa: {r4}")
    ok &= r4 == shutil.which("ffmpeg")
    bot.RAIZ = raiz_original

    # 5) detecção de plataforma do baixar_ffmpeg.py
    original_sistema, original_maquina = _platform.system, _platform.machine
    casos = [
        ("Windows", "AMD64", "win64", ".zip"),
        ("Windows", "ARM64", "winarm64", ".zip"),
        ("Linux", "x86_64", "linux64", ".tar.xz"),
        ("Linux", "aarch64", "linuxarm64", ".tar.xz"),
        ("Darwin", "x86_64", "osx64", ".zip"),
    ]
    for sistema, maquina, alvo, ext in casos:
        _platform.system = lambda s=sistema: s
        _platform.machine = lambda m=maquina: m
        obtido = baixar_ffmpeg.alvo_plataforma()
        casou = obtido == (alvo, ext)
        ok &= casou
        print(f"  {'✅' if casou else '❌'} {sistema}/{maquina:8} -> {obtido[0]}{obtido[1]}")
    _platform.system = lambda: "Windows"
    _platform.machine = lambda: "sparc"
    try:
        baixar_ffmpeg.alvo_plataforma()
        ok = False
        print("  ❌ plataforma desconhecida deveria dar erro")
    except SystemExit:
        print("  ✅ plataforma desconhecida -> erro explicativo")
    _platform.system, _platform.machine = original_sistema, original_maquina

    # 6) o instalador reconhece o ffmpeg que já existe
    encontrado, onde = baixar_ffmpeg.situacao_atual(tmp / "vazio")
    print(f"  --conferir acha o ffmpeg do sistema: {encontrado} ({onde})")
    ok &= encontrado is not None and baixar_ffmpeg.versao_do(encontrado).startswith("ffmpeg version")

    # 7) binário falso é reconhecido como portátil válido
    encontrado2, onde2 = baixar_ffmpeg.situacao_atual(bin_falso / "bin")
    print(f"  --conferir acha o portátil da pasta bin/: {encontrado2} ({onde2})")
    ok &= str(encontrado2) == str(exe) and "portátil" in onde2

    print("RESULTADO:", "OK ✅" if ok else "FALHOU ❌")
    return ok


def main() -> int:
    import logging

    logging.basicConfig(level=logging.WARNING)
    with tempfile.TemporaryDirectory(prefix="hl-testes-") as d:
        tmp = Path(d)
        ok1 = teste_conversao(tmp)
        ok2 = teste_download(tmp)
        ok3 = teste_regras(tmp)
        ok4 = teste_ponta_a_ponta(tmp)
        ok5 = teste_youtube(tmp)
        ok6 = teste_ffmpeg(tmp)
    print()
    print("=" * 62)
    print("TODOS OS TESTES:", "OK ✅" if all((ok1, ok2, ok3, ok4, ok5, ok6)) else "HOUVE FALHA ❌")
    print("=" * 62)
    return 0 if all((ok1, ok2, ok3, ok4, ok5, ok6)) else 1


if __name__ == "__main__":
    sys.exit(main())
