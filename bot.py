#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Bot do Telegram x Holyrics
==========================

Recebe arquivos enviados no Telegram e salva cada um na pasta correspondente
do Holyrics (configurável no config.json). PDF é convertido para PowerPoint
(.pptx), que é o formato de apresentação aceito pelo Holyrics.

- Mostra o progresso do download na MESMA mensagem (sem spam) e informa no
  final se houve conversão e onde o arquivo foi salvo.
- Roda como serviço no Windows (Task Scheduler) e no Linux (systemd).

Uso:  python bot.py [--config config.json] [--verificar]
"""

from __future__ import annotations

import argparse
import asyncio
import base64
import html
import json
import logging
import logging.handlers
import os
import re
import secrets
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
from pathlib import Path
from urllib.parse import unquote, urlparse

import httpx
from telegram import BotCommand, InlineKeyboardButton, InlineKeyboardMarkup, Update
from minio_holyrics import MinioStore
from telegram.constants import ChatAction, ParseMode
from telegram.error import RetryAfter
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

LOG = logging.getLogger("holyrics-bot")
RAIZ = Path(__file__).resolve().parent
BASE_API = "https://api.telegram.org"  # trocável nos testes

# ---------------------------------------------------------------------------
# Categorias de arquivo -> chave da pasta no config.json
# ---------------------------------------------------------------------------

CATEGORIAS: dict[str, set[str]] = {
    "audio": {
        ".mp3", ".wav", ".m4a", ".aac", ".ogg", ".oga", ".opus", ".flac",
        ".wma", ".amr", ".mp2", ".aiff", ".aif", ".ac3", ".mka", ".weba",
    },
    "video": {
        ".mp4", ".mkv", ".avi", ".mov", ".wmv", ".webm", ".m4v", ".mpg",
        ".mpeg", ".3gp", ".flv", ".ts", ".mts", ".vob", ".ogv",
    },
    "imagem": {
        ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".tif", ".tiff",
        ".heic", ".heif", ".svg", ".ico", ".jfif", ".avif",
    },
    "apresentacao": {
        ".pptx", ".ppt", ".ppsx", ".pps", ".odp", ".key", ".pptm",
    },
    "documento": {
        ".txt", ".doc", ".docx", ".rtf", ".odt", ".csv", ".xls", ".xlsx",
        ".ods", ".md", ".sub", ".srt", ".vtt",
    },
}

MIME_EXT = {
    "image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp",
    "image/gif": ".gif", "image/bmp": ".bmp", "image/tiff": ".tiff",
    "audio/mpeg": ".mp3", "audio/mp4": ".m4a", "audio/ogg": ".ogg",
    "audio/opus": ".opus", "audio/wav": ".wav", "audio/x-wav": ".wav",
    "audio/flac": ".flac", "audio/aac": ".aac", "audio/amr": ".amr",
    "video/mp4": ".mp4", "video/quicktime": ".mov", "video/x-matroska": ".mkv",
    "video/webm": ".webm", "video/x-msvideo": ".avi",
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
}


def categoria_de(nome: str) -> str:
    """Descobre a categoria (pasta) pela extensão do arquivo."""
    ext = Path(nome).suffix.lower()
    for cat, exts in CATEGORIAS.items():
        if ext in exts:
            return cat
    return "outros"


def eh_pdf(nome: str) -> bool:
    return Path(nome).suffix.lower() == ".pdf"


# ---------------------------------------------------------------------------
# Configuração
# ---------------------------------------------------------------------------

class Config:
    def __init__(self, caminho: Path):
        self.caminho = caminho
        try:
            bruto = json.loads(caminho.read_text(encoding="utf-8"))
        except FileNotFoundError:
            sys.exit(
                f"ERRO: config não encontrado em {caminho}\n"
                f"Copie config.example.json para config.json e preencha o token do bot."
            )
        except json.JSONDecodeError as e:
            sys.exit(f"ERRO: config.json inválido: {e}")

        tg = bruto.get("telegram") or {}
        self.token: str = (tg.get("token") or "").strip()
        self.usuarios: set[int] = {int(u) for u in (tg.get("usuarios_autorizados") or [])}
        self.aceitar_grupos: bool = bool(tg.get("aceitar_grupos", False))
        self.admin: int | None = int(tg["admin"]) if tg.get("admin") else None

        # Servidor local do Telegram Bot API (remove o limite de 20 MB).
        self.api_id: int = int(tg.get("api_id") or 0)
        self.api_hash: str = str(tg.get("api_hash") or "").strip()
        self.base_url: str = str(
            tg.get("base_url") or "https://api.telegram.org/bot"
        ).rstrip("/")
        self.base_file_url: str = str(
            tg.get("base_file_url") or "https://api.telegram.org/file/bot"
        ).rstrip("/")
        self.local_mode: bool = bool(tg.get("local_mode", False))

        self.pastas: dict[str, Path] = {
            cat: self._resolver(caminho, cam)
            for cat, cam in (bruto.get("pastas") or {}).items()
        }

        pdf = bruto.get("pdf") or {}
        self.pdf_converter: bool = bool(pdf.get("converter_para_pptx", True))
        self.pdf_dpi: int = int(pdf.get("dpi", 200))
        self.pdf_largura_max: int = int(pdf.get("largura_max_px", 2400))
        self.pdf_qualidade: int = int(pdf.get("qualidade_jpeg", 90))
        self.pdf_manter_original: bool = bool(pdf.get("manter_pdf_original", False))

        # Conversão de mídia para os formatos que o Holyrics prefere (mp3/mp4).
        cv = bruto.get("converter") or {}
        self.converter_perguntar: bool = bool(cv.get("perguntar", True))
        self.converter_audio_kbps: int = int(cv.get("audio_bitrate_kbps", 192))

        dl = bruto.get("download") or {}
        self.progresso_intervalo: float = float(dl.get("intervalo_progresso_segundos", 2))
        self.max_mb: int = int(dl.get("tamanho_maximo_mb", 20))

        # Baixar arquivos a partir de um link colado no chat.
        lk = bruto.get("links") or {}
        self.links_ativo: bool = bool(lk.get("ativo", True))
        self.links_max_mb: int = int(lk.get("tamanho_maximo_mb", 0))  # 0 = sem limite
        self.links_bloqueados: list[str] = [
            str(h).lower() for h in (lk.get("hosts_bloqueados") or [])
        ]

        # MinIO: usado para receber arquivos maiores que o limite do Telegram.
        mn = bruto.get("minio") or {}
        self.minio_ativo: bool = bool(mn.get("ativo", False))
        self.minio_endpoint: str = str(mn.get("endpoint") or "").strip()
        self.minio_access: str = str(mn.get("access_key") or "").strip()
        self.minio_secret: str = str(mn.get("secret_key") or "").strip()
        self.minio_bucket: str = str(mn.get("bucket") or "holyrics-entrada").strip()
        self.minio_secure: bool = bool(mn.get("secure", True))
        self.minio_regiao: str = str(mn.get("regiao") or "us-east-1").strip()
        self.minio_pagina_upload: str = str(mn.get("url_pagina_upload") or "").strip()
        self.minio_bucket_publico: str = str(mn.get("bucket_publico") or "").strip()
        self.minio_prefixo: str = str(mn.get("prefixo") or "entrada/").strip()
        self.minio_expira_min: int = int(mn.get("expira_min", 120))
        self.minio_intervalo_s: int = int(mn.get("intervalo_verificacao_s", 10))
        self.minio_espera_min: int = int(mn.get("espera_max_min", 120))
        self.minio_max_mb: int = int(mn.get("tamanho_maximo_mb", 0))  # 0 = sem limite

        yt = bruto.get("youtube") or {}
        self.yt_ativo: bool = bool(yt.get("ativo", True))
        self.yt_altura_maxima: int = int(yt.get("altura_maxima", 1080))
        self.yt_duracao_max_min: int = int(yt.get("duracao_maxima_min", 120))
        self.yt_mp3_kbps: int = int(yt.get("mp3_bitrate_kbps", 192))
        self.yt_ffmpeg: str = (yt.get("ffmpeg") or "").strip()
        self.yt_cookies: str = (yt.get("cookies") or "").strip()

        lg = bruto.get("log") or {}
        self.log_arquivo: str = lg.get("arquivo", "bot.log")
        self.log_nivel: str = (lg.get("nivel", "INFO") or "INFO").upper()

        if not self.token:
            sys.exit("ERRO: 'telegram.token' está vazio no config.json")

    @staticmethod
    def _resolver(caminho_config: Path, valor: str) -> Path:
        """Expande {USERPROFILE}, %VARS%, ~ e caminhos relativos ao config.json."""
        s = str(valor)
        s = s.replace("{USERPROFILE}", str(Path.home()))
        s = s.replace("{DOCUMENTS}", str(Path.home() / "Documents"))
        s = s.replace("{HOME}", str(Path.home()))
        s = os.path.expandvars(os.path.expanduser(s))
        p = Path(s)
        if not p.is_absolute():
            p = (caminho_config.parent / p).resolve()
        return p

    def minio_configurado(self) -> bool:
        return bool(
            self.minio_ativo
            and self.minio_endpoint
            and self.minio_access
            and self.minio_secret
        )


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def fmt_bytes(n: int) -> str:
    if n <= 0:
        return "0 B"
    for unidade in ("B", "KB", "MB", "GB"):
        if n < 1024 or unidade == "GB":
            if unidade == "B":
                return f"{int(n)} B"
            return f"{n:.1f} {unidade}".replace(".", ",")
        n /= 1024.0
    return f"{n:.1f} GB"


def nome_seguro(nome: str) -> str:
    """Remove diretórios e caracteres inválidos do nome do arquivo."""
    nome = unicodedata.normalize("NFC", nome or "")
    nome = nome.replace("\\", "/").split("/")[-1]
    nome = re.sub(r'[<>:"|?*\x00-\x1f]', "_", nome).strip(" .")
    return nome or "arquivo"


def nome_livre(destino: Path) -> Path:
    """Se o arquivo já existe, gera 'nome (2).ext', 'nome (3).ext'..."""
    if not destino.exists():
        return destino
    base, suf, i = destino.stem, destino.suffix, 2
    while True:
        alt = destino.with_name(f"{base} ({i}){suf}")
        if not alt.exists():
            return alt
        i += 1


def estado_pasta(p: Path) -> tuple[str, str]:
    """Diz se a pasta existe / será criada / é inválida."""
    if p.is_dir():
        return "✅", "existe"
    ancestral = p
    while not ancestral.exists() and ancestral != ancestral.parent:
        ancestral = ancestral.parent
    if ancestral.is_dir() and os.access(ancestral, os.W_OK):
        return "⚠️", f"não existe (o bot cria — ancestral: {ancestral})"
    return "❌", "caminho inválido ou sem permissão de escrita"


def template_estado(titulo_html: str, tamanho: int) -> str:
    return f"📥 Recebi {titulo_html} ({fmt_bytes(tamanho)})\nBaixando… 0%"


def fmt_duracao(segundos: float | int | None) -> str:
    if not segundos:
        return "?"
    segundos = int(segundos)
    h, resto = divmod(segundos, 3600)
    m, s = divmod(resto, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def caminho_ffmpeg(cfg: "Config") -> str | None:
    """Descobre o ffmpeg: youtube.ffmpeg -> pasta bin/ do bot (portátil) -> PATH."""
    if cfg.yt_ffmpeg:
        if Path(cfg.yt_ffmpeg).exists():
            return cfg.yt_ffmpeg
        LOG.warning("youtube.ffmpeg aponta para um arquivo que não existe (%s) — tentando os próximos", cfg.yt_ffmpeg)
    nome = "ffmpeg.exe" if os.name == "nt" else "ffmpeg"
    local = RAIZ / "bin" / nome
    if local.exists():
        return str(local)
    return shutil.which("ffmpeg")


def ffmpeg_disponivel(cfg: "Config") -> bool:
    return caminho_ffmpeg(cfg) is not None


# ---------------------------------------------------------------------------
# Download com progresso (mesma mensagem, sem spam)
# ---------------------------------------------------------------------------

class Progresso:
    """Atualiza a mensagem de status no máximo a cada X segundos / Y pontos."""

    def __init__(self, mensagem, titulo_html: str, total: int, intervalo: float):
        self.mensagem = mensagem
        self.titulo = titulo_html
        self.total = total
        self.intervalo = intervalo
        self.ultimo_t = time.monotonic()
        self.ultimo_pct = -1

    def _texto(self, feito: int) -> str:
        if self.total:
            pct = min(99, int(feito * 100 / self.total))
            return (
                f"📥 Recebi {self.titulo} ({fmt_bytes(self.total)})\n"
                f"Baixando… {pct}%  ({fmt_bytes(feito)} de {fmt_bytes(self.total)})"
            )
        return (
            f"📥 Recebi {self.titulo}\nBaixando… {fmt_bytes(feito)}"
        )

    async def __call__(self, feito: int) -> None:
        if self.total:
            pct = int(feito * 100 / self.total)
            if pct == self.ultimo_pct or pct < 5:
                return
            self.ultimo_pct = pct
        if time.monotonic() - self.ultimo_t < self.intervalo:
            return
        self.ultimo_t = time.monotonic()
        texto = self._texto(feito)
        try:
            await self.mensagem.edit_text(texto, parse_mode=ParseMode.HTML)
        except RetryAfter as e:
            LOG.debug("Rate limit do Telegram, aguardando %ss", e.retry_after)
            self.ultimo_t += float(e.retry_after)
        except Exception as e:  # mensagem apagada, rede etc. — não interrompe o download
            LOG.debug("Falha ao atualizar progresso: %s", e)


async def baixar(url: str, destino: Path, progresso: Progresso) -> int:
    """Baixa em streaming, reportando o progresso a cada bloco."""
    parcial = destino.with_name(destino.name + ".parcial")
    feito = 0
    timeout = httpx.Timeout(connect=30.0, read=60.0, write=60.0, pool=30.0)
    async with httpx.AsyncClient(timeout=timeout, follow_redirects=True) as cliente:
        async with cliente.stream("GET", url) as resposta:
            resposta.raise_for_status()
            with open(parcial, "wb") as saida:
                async for bloco in resposta.aiter_bytes(512 * 1024):
                    saida.write(bloco)
                    feito += len(bloco)
                    await progresso(feito)
    os.replace(parcial, destino)
    return feito


# ---------------------------------------------------------------------------
# PDF -> PowerPoint
# ---------------------------------------------------------------------------

def pdf_para_pptx(origem: Path, saida: Path, cfg: Config) -> int:
    """Converte cada página do PDF em um slide (formato 16:9) do PowerPoint."""
    import pypdfium2 as pdfium
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    LARG_IN, ALT_IN = 13.333, 7.5  # slide widescreen padrão
    doc = pdfium.PdfDocument(str(origem))
    try:
        total = len(doc)
        if total == 0:
            raise ValueError("PDF sem páginas")

        prs = Presentation()
        prs.slide_width = Inches(LARG_IN)
        prs.slide_height = Inches(ALT_IN)
        layout_branco = prs.slide_layouts[6]

        with tempfile.TemporaryDirectory(prefix="hl-pdf-") as tmp:
            for i in range(total):
                pagina = doc[i]
                largura_pt, altura_pt = pagina.get_size()
                bmp = pagina.render(scale=cfg.pdf_dpi / 72)
                img = bmp.to_pil()

                # downscale se passar do limite, e garante RGB (fundo branco)
                if img.width > cfg.pdf_largura_max:
                    nova_alt = round(img.height * cfg.pdf_largura_max / img.width)
                    img = img.resize((cfg.pdf_largura_max, nova_alt), Image.LANCZOS)
                if img.mode in ("RGBA", "LA", "P"):
                    img = img.convert("RGBA")
                    fundo = Image.new("RGB", img.size, "white")
                    fundo.paste(img, mask=img.split()[-1])
                    img = fundo
                elif img.mode != "RGB":
                    img = img.convert("RGB")

                jpg = Path(tmp) / f"pagina-{i + 1:04d}.jpg"
                img.save(jpg, "JPEG", quality=cfg.pdf_qualidade, optimize=True)

                # encaixa a página dentro do slide 16:9, centralizada
                razao = largura_pt / altura_pt if altura_pt else (LARG_IN / ALT_IN)
                larg = LARG_IN
                alt = larg / razao
                if alt > ALT_IN:
                    alt = ALT_IN
                    larg = alt * razao

                slide = prs.slides.add_slide(layout_branco)
                slide.shapes.add_picture(
                    str(jpg),
                    Inches((LARG_IN - larg) / 2),
                    Inches((ALT_IN - alt) / 2),
                    width=Inches(larg),
                    height=Inches(alt),
                )
        prs.save(str(saida))
        return total
    finally:
        doc.close()


# ---------------------------------------------------------------------------
# YouTube (yt-dlp)
# ---------------------------------------------------------------------------

RE_YOUTUBE = re.compile(
    r"(?:https?://)?(?:(?:www|m|music)\.)?"
    r"(?:youtube\.com/(?:watch\?[^\s]*?\bv=|shorts/|live/|embed/|v/)|youtu\.be/)"
    r"([A-Za-z0-9_-]{11})"
)


def achar_videos_youtube(texto: str) -> list[str]:
    """Extrai links de vídeo do YouTube de um texto (sem repetir)."""
    vistos: set[str] = set()
    urls: list[str] = []
    for achado in RE_YOUTUBE.finditer(texto or ""):
        vid = achado.group(1)
        if vid not in vistos:
            vistos.add(vid)
            urls.append(f"https://www.youtube.com/watch?v={vid}")
    return urls


def opcoes_ytdlp(cfg: Config) -> dict:
    opts: dict = {"quiet": True, "no_warnings": True, "noplaylist": True}
    if cfg.yt_cookies:
        opts["cookiefile"] = cfg.yt_cookies
    ffmpeg = caminho_ffmpeg(cfg)
    if ffmpeg:
        opts["ffmpeg_location"] = ffmpeg
    return opts


def info_youtube(url: str, cfg: Config) -> dict:
    """Consulta os dados do vídeo (título, duração, formatos disponíveis)."""
    import yt_dlp

    opts = opcoes_ytdlp(cfg)
    opts["skip_download"] = True
    with yt_dlp.YoutubeDL(opts) as ydl:
        return ydl.extract_info(url, download=False)


def _tamanho_estimado(f: dict, duracao: float | None) -> int:
    tam = f.get("filesize") or f.get("filesize_approx")
    if not tam and f.get("tbr") and duracao:
        tam = float(f["tbr"]) * 1000 / 8 * duracao
    return int(tam or 0)


def _formatos_audio(formatos: list[dict]) -> list[dict]:
    return [f for f in formatos if f.get("vcodec") == "none" and f.get("acodec") not in (None, "none")]


def _formatos_video(formatos: list[dict]) -> list[dict]:
    return [f for f in formatos if f.get("height") and f.get("vcodec") not in (None, "none")]


def estimativa_video(formatos: list[dict], altura: int, duracao: float | None) -> int:
    videos = [f for f in _formatos_video(formatos) if int(f["height"]) <= altura]
    if not videos:
        return 0
    audios = _formatos_audio(formatos)
    total = _tamanho_estimado(max(videos, key=lambda f: int(f["height"])), duracao)
    if audios:
        total += _tamanho_estimado(max(audios, key=lambda f: _tamanho_estimado(f, duracao)), duracao)
    return total


def estimativa_audio(formatos: list[dict], duracao: float | None) -> int:
    return max((_tamanho_estimado(f, duracao) for f in _formatos_audio(formatos)), default=0)


def montar_opcoes(info: dict, cfg: Config) -> tuple[str, list[list[tuple[str, int]]], list[dict]]:
    """Monta o texto e os botões (2 de vídeo + 2 de áudio) para o vídeo."""
    titulo = info.get("title") or "vídeo"
    canal = info.get("uploader") or info.get("channel") or "?"
    duracao = info.get("duration")
    formatos = info.get("formats") or []
    com_ffmpeg = ffmpeg_disponivel(cfg)

    alturas = sorted(
        {int(f["height"]) for f in _formatos_video(formatos) if int(f["height"]) <= cfg.yt_altura_maxima},
        reverse=True,
    )
    escolhidas: list[int | None] = []
    if alturas:
        escolhidas.append(alturas[0])
        menor = next((h for h in alturas[1:] if h <= alturas[0] * 0.75), None)
        escolhidas.append(menor if menor else (alturas[1] if len(alturas) > 1 else None))
    else:
        escolhidas = [None]

    selecoes: list[dict] = []
    linha_video: list[tuple[str, int]] = []
    for altura in escolhidas:
        if altura is None:
            if not linha_video:
                formato = "bestvideo+bestaudio/best" if com_ffmpeg else "best"
                rotulo = "Melhor disponível"
                estimativa = 0
            else:
                continue
        else:
            altura = int(altura)
            if com_ffmpeg:
                formato = (
                    f"bestvideo[height<={altura}][ext=mp4]+bestaudio[ext=m4a]/"
                    f"bestvideo[height<={altura}]+bestaudio/best[height<={altura}]"
                )
            else:
                formato = f"best[height<={altura}][ext=mp4]/best[height<={altura}]"
            rotulo = f"{altura}p"
            estimativa = estimativa_video(formatos, altura, duracao)
        indice = len(selecoes)
        selecoes.append({"formato": formato, "mp3": False, "rotulo": rotulo, "tipo": "video"})
        etiqueta = f"🎬 {rotulo}" + (f" · ~{fmt_bytes(estimativa)}" if estimativa else "")
        linha_video.append((etiqueta, indice))

    estim_audio = estimativa_audio(formatos, duracao)
    if com_ffmpeg:
        # O Holyrics trabalha melhor com MP3, então todo áudio vira MP3.
        pares_audio = [
            ({"formato": "bestaudio/best", "mp3": True, "kbps": cfg.yt_mp3_kbps,
              "rotulo": f"MP3 {cfg.yt_mp3_kbps} kbps", "tipo": "audio"},
             f"🎵 MP3 {cfg.yt_mp3_kbps}k"),
            ({"formato": "bestaudio/best", "mp3": True, "kbps": 128,
              "rotulo": "MP3 128 kbps", "tipo": "audio"},
             "🎵 MP3 128k (leve)"),
        ]
    else:
        pares_audio = [
            ({"formato": "bestaudio[ext=m4a]/bestaudio", "mp3": False, "rotulo": "Melhor áudio (m4a)", "tipo": "audio"},
             "🎵 Melhor áudio"),
            ({"formato": "bestaudio[abr<=64]/bestaudio", "mp3": False, "rotulo": "Áudio leve (m4a)", "tipo": "audio"},
             f"🎵 Leve {fmt_bytes(estim_audio // 4) if estim_audio else '~64k'}"),
        ]
    linha_audio: list[tuple[str, int]] = []
    for selecao, etiqueta in pares_audio:
        indice = len(selecoes)
        selecoes.append(selecao)
        if estim_audio and selecao == pares_audio[0][0]:
            etiqueta += f" · ~{fmt_bytes(estim_audio)}"
        linha_audio.append((etiqueta, indice))

    linhas = [
        f"🎬 <b>{html.escape(titulo)}</b>",
        f"👤 {html.escape(str(canal))} · ⏱ {fmt_duracao(duracao)}",
        "Escolha o formato para baixar:",
    ]
    if not com_ffmpeg:
        linhas.append("⚠️ ffmpeg não encontrado: sem MP3 e sem juntar vídeo+áudio.")

    return "\n".join(linhas), [linha_video, linha_audio], selecoes


def localizar_arquivo_baixado(pasta: Path, info: dict) -> Path | None:
    """Descobre o arquivo final gerado pelo yt-dlp dentro da pasta temporária."""
    candidatos = [
        p for p in pasta.rglob("*")
        if p.is_file()
        and not p.name.endswith((".part", ".ytdl", ".temp"))
        and not re.search(r"\.f\d+\.\w+$", p.name)
    ]
    if not candidatos:
        return None
    for pedido in (info.get("requested_downloads") or []):
        caminho = Path(pedido.get("filepath") or "")
        if caminho and caminho.exists():
            return caminho
        if caminho:
            for c in candidatos:
                if c.with_suffix("") == caminho.with_suffix(""):
                    return c
    return max(candidatos, key=lambda p: p.stat().st_size)


def baixar_youtube(estado: dict, url: str, pasta: Path, selecao: dict, cfg: Config) -> tuple[Path, dict]:
    """Baixa (e converte, se MP3) o vídeo. Roda em thread; informa o progresso via `estado`."""
    import yt_dlp

    tmp = Path(tempfile.mkdtemp(prefix="yt-bot-"))

    def hook(d):
        status = d.get("status")
        if status == "downloading":
            estado.update(
                fase="baixando",
                feito=d.get("downloaded_bytes") or 0,
                total=d.get("total_bytes") or d.get("total_bytes_estimate") or 0,
                vel=d.get("speed") or 0,
                eta=d.get("eta") or 0,
            )
        elif status == "finished":
            estado.update(fase="processando", vel=0, eta=0)

    def hook_pp(d):
        if d.get("status") == "started":
            estado.update(fase="convertendo", pp=d.get("postprocessor"))

    opts = opcoes_ytdlp(cfg)
    opts.update({
        "outtmpl": str(tmp / "%(title).120B.%(ext)s"),
        "format": selecao["formato"],
        "progress_hooks": [hook],
        "postprocessor_hooks": [hook_pp],
        "noprogress": True,
        "retries": 5,
        "fragment_retries": 5,
        "concurrent_fragment_downloads": 4,
        "trim_file_name": 120,
        "windowsfilenames": True,
    })
    if selecao["tipo"] == "video":
        opts["merge_output_format"] = "mp4"
    if selecao.get("mp3"):
        opts["postprocessors"] = [{
            "key": "FFmpegExtractAudio",
            "preferredcodec": "mp3",
            "preferredquality": str(selecao.get("kbps") or cfg.yt_mp3_kbps),
        }]

    try:
        with yt_dlp.YoutubeDL(opts) as ydl:
            info = ydl.extract_info(url, download=True)
        estado["fase"] = "salvando"
        origem = localizar_arquivo_baixado(tmp, info)
        if origem is None:
            raise RuntimeError("o download terminou mas nenhum arquivo foi encontrado")
        pasta.mkdir(parents=True, exist_ok=True)
        destino = nome_livre(pasta / nome_seguro(origem.name))
        shutil.move(str(origem), str(destino))
        return destino, info
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def texto_progresso_yt(estado: dict, titulo_html: str, rotulo: str) -> str:
    fase = estado.get("fase", "iniciando")
    if fase == "baixando":
        feito = estado.get("feito") or 0
        total = estado.get("total") or 0
        pct = int(feito * 100 / total) if total else 0
        partes = [f"{pct}%", fmt_bytes(feito) + (f" de {fmt_bytes(total)}" if total else "")]
        if estado.get("vel"):
            partes.append(f"{fmt_bytes(int(estado['vel']))}/s")
        if estado.get("eta"):
            partes.append(f"faltam {fmt_duracao(estado['eta'])}")
        return f"⬇️ Baixando <b>{titulo_html}</b> ({rotulo})\n" + " · ".join(partes)
    if fase == "convertendo":
        pp = (estado.get("pp") or "").lower()
        if "extractaudio" in pp:
            return f"🎧 Convertendo para {rotulo}…\n<b>{titulo_html}</b>"
        if "merge" in pp:
            return f"🎬 Juntando vídeo + áudio…\n<b>{titulo_html}</b>"
        return f"⚙️ Processando <b>{titulo_html}</b>…"
    if fase == "processando":
        return f"⚙️ Finalizando o download de <b>{titulo_html}</b>…"
    if fase == "salvando":
        return f"📦 Salvando <b>{titulo_html}</b>…"
    return f"⏳ Preparando <b>{titulo_html}</b>…"


async def editar(mensagem, texto: str, **extra) -> None:
    """Edita a mensagem de status, ignorando erros de rede/rate limit."""
    try:
        await mensagem.edit_text(texto, parse_mode=ParseMode.HTML, **extra)
    except RetryAfter as e:
        await asyncio.sleep(float(e.retry_after))
    except Exception as e:
        LOG.debug("Falha ao editar a mensagem: %s", e)


async def rodar_com_progresso(mensagem, estado: dict, titulo_html: str, rotulo: str,
                              funcao, intervalo: float, *args):
    """Roda uma função bloqueante informando o progresso na mensagem."""
    parar = asyncio.Event()

    async def vigia():
        ultimo = None
        while not parar.is_set():
            texto = texto_progresso_yt(estado, titulo_html, rotulo)
            if texto != ultimo:
                ultimo = texto
                await editar(mensagem, texto)
            try:
                await asyncio.wait_for(parar.wait(), intervalo)
            except asyncio.TimeoutError:
                pass

    tarefa = asyncio.create_task(vigia())
    try:
        return await asyncio.to_thread(funcao, estado, *args)
    finally:
        parar.set()
        await tarefa


def dica_erro_youtube(e: Exception) -> str:
    texto = str(e)
    baixo = texto.lower()
    if "sign in to confirm" in baixo or "not a bot" in baixo or "cookies" in baixo:
        return (
            "\n\n👉 O YouTube pediu verificação de robô. Exporte os cookies do navegador para um "
            "arquivo no formato Netscape e aponte em <code>youtube.cookies</code> no config.json "
            "(ou tente de novo mais tarde)."
        )
    if "ffmpeg" in baixo:
        return "\n\n👉 O ffmpeg não foi encontrado. Instale o ffmpeg ou aponte o caminho em <code>youtube.ffmpeg</code>."
    if "unavailable" in baixo or "private" in baixo:
        return "\n\n👉 O vídeo parece privado ou indisponível."
    if "unsupported url" in baixo:
        return "\n\n👉 Esse link não é de um vídeo do YouTube."
    return ""


async def tratar_links(update: Update, context: ContextTypes.DEFAULT_TYPE, urls: list[str]) -> None:
    cfg: Config = context.bot_data["cfg"]
    msg = update.effective_message
    if not cfg.yt_ativo:
        await msg.reply_text("🔗 O download do YouTube está desativado no config.json.")
        return

    pendentes: dict = context.bot_data.setdefault("yt_pendentes", {})
    for url in urls[:3]:
        aviso = await msg.reply_text("🔎 Verificando o vídeo…")
        try:
            info = await asyncio.to_thread(info_youtube, url, cfg)
        except Exception as e:
            LOG.warning("Falha ao consultar %s: %s", url, e)
            await editar(aviso, f"❌ Não consegui verificar esse vídeo.\n{html.escape(str(e))}{dica_erro_youtube(e)}")
            continue

        duracao = info.get("duration") or 0
        if cfg.yt_duracao_max_min and duracao and duracao > cfg.yt_duracao_max_min * 60:
            await editar(
                aviso,
                f"⏱ <b>{html.escape(info.get('title') or url)}</b> tem {fmt_duracao(duracao)} e passa "
                f"do limite de {cfg.yt_duracao_max_min} min configurado.",
            )
            continue

        texto, linhas_botoes, selecoes = montar_opcoes(info, cfg)
        token = secrets.token_hex(4)
        if len(pendentes) > 200:
            pendentes.clear()
        pendentes[token] = {
            "url": url,
            "info": info,
            "user": update.effective_user.id,
            "chat": msg.chat_id,
            "selecoes": selecoes,
            "quando": time.time(),
        }
        teclado = InlineKeyboardMarkup([
            [InlineKeyboardButton(etiqueta, callback_data=f"yt|{token}|{indice}") for etiqueta, indice in linha]
            for linha in linhas_botoes if linha
        ])
        LOG.info("Opções enviadas para %s (token %s)", url, token)
        await editar(aviso, texto, reply_markup=teclado)


async def botao_youtube(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    cfg: Config = context.bot_data["cfg"]
    consulta = update.callback_query
    partes = (consulta.data or "").split("|")
    pendentes: dict = context.bot_data.setdefault("yt_pendentes", {})
    pendente = pendentes.pop(partes[1], None) if len(partes) == 3 else None
    mensagem = consulta.message

    if pendente is None:
        await consulta.answer("⌛ Esta escolha expirou. Mande o link de novo.", show_alert=True)
        await editar(mensagem, "⌛ Opção expirada — mande o link do YouTube novamente.")
        return

    if consulta.from_user.id != pendente["user"]:
        await consulta.answer("Esse download é de outro usuário. Mande o link você mesmo.", show_alert=True)
        return

    selecao = pendente["selecoes"][int(partes[2])]
    await consulta.answer()
    try:
        await mensagem.edit_reply_markup(None)
    except Exception:
        pass

    titulo = pendente["info"].get("title") or "vídeo"
    titulo_html = html.escape(titulo)
    rotulo = selecao["rotulo"]
    chave_pasta = "audio" if selecao["tipo"] == "audio" else "video"
    pasta = cfg.pastas.get(chave_pasta) or cfg.pastas.get("outros")
    if pasta is None:
        await editar(mensagem, f"❌ Nenhuma pasta configurada para <b>{chave_pasta}</b>.")
        return

    LOG.info("Baixando %s em %s (%s)", pendente["url"], rotulo, pasta)
    estado: dict = {"fase": "iniciando", "feito": 0, "total": 0}
    try:
        destino, _ = await rodar_com_progresso(
            mensagem, estado, titulo_html, rotulo, baixar_youtube,
            cfg.progresso_intervalo, pendente["url"], pasta, selecao, cfg,
        )
    except Exception as e:
        LOG.exception("Falha no download de %s", pendente["url"])
        await editar(mensagem, f"❌ Falha ao baixar <b>{titulo_html}</b>:\n{html.escape(str(e))}{dica_erro_youtube(e)}")
        return

    duracao = pendente["info"].get("duration")
    icone = "🎵" if selecao["tipo"] == "audio" else "🎬"
    extra = f" · ⏱ {fmt_duracao(duracao)}" if duracao else ""
    LOG.info("Baixado: %s", destino)
    await editar(
        mensagem,
        f"✅ <b>{titulo_html}</b>\n"
        f"{icone} {rotulo}\n"
        f"📁 Salvo em:\n<code>{html.escape(str(destino))}</code>\n"
        f"📦 {fmt_bytes(destino.stat().st_size)}{extra}",
    )


# ---------------------------------------------------------------------------
# Bot
# ---------------------------------------------------------------------------

def autorizado(update: Update, cfg: Config) -> bool:
    usuario = update.effective_user
    if usuario is None:
        return False
    if not cfg.usuarios:
        return True  # config sem lista = liberado (útil para testar)
    if usuario.id in cfg.usuarios:
        return True
    if cfg.admin and usuario.id == cfg.admin:
        return True
    return False


def extrair_arquivo(msg) -> tuple[str, str, int] | None:
    """Retorna (file_id, nome, tamanho) do arquivo da mensagem, se houver."""
    agora = time.strftime("%Y%m%d-%H%M%S")

    def ext_mime(mime: str | None, padrao: str) -> str:
        return MIME_EXT.get((mime or "").lower(), padrao)

    if msg.document:
        nome = msg.document.file_name or f"documento-{agora}{ext_mime(msg.document.mime_type, '.bin')}"
        return msg.document.file_id, nome, msg.document.file_size or 0
    if msg.audio:
        nome = msg.audio.file_name or f"audio-{agora}{ext_mime(msg.audio.mime_type, '.mp3')}"
        return msg.audio.file_id, nome, msg.audio.file_size or 0
    if msg.video:
        nome = msg.video.file_name or f"video-{agora}{ext_mime(msg.video.mime_type, '.mp4')}"
        return msg.video.file_id, nome, msg.video.file_size or 0
    if msg.voice:
        return msg.voice.file_id, f"voz-{agora}.ogg", msg.voice.file_size or 0
    if msg.video_note:
        return msg.video_note.file_id, f"video-nota-{agora}.mp4", msg.video_note.file_size or 0
    if msg.animation:
        nome = msg.animation.file_name or f"animacao-{agora}.mp4"
        return msg.animation.file_id, nome, msg.animation.file_size or 0
    if msg.photo:
        maior = msg.photo[-1]
        return maior.file_id, f"foto-{agora}.jpg", maior.file_size or 0
    if msg.sticker:
        suf = ".webp" if not msg.sticker.is_animated else ".tgs"
        return msg.sticker.file_id, f"sticker-{agora}{suf}", msg.sticker.file_size or 0
    return None


async def cmd_start(update: Update, cfg: Config) -> None:
    ajuda = (
        "<b>Bot do Holyrics</b>\n\n"
        "Me envie um arquivo e eu salvo na pasta correspondente:\n"
        "• áudio → pasta de áudios\n"
        "• vídeo → pasta de vídeos\n"
        "• imagem → pasta de imagens\n"
        "• PDF → convertido para PowerPoint e salvo na pasta de apresentações\n"
        "• PowerPoint/documentos → pasta correspondente\n\n"
        "Ou me mande um <b>link do YouTube</b>: eu mostro o título, a duração e ofereço\n"
        "duas opções de vídeo e duas de áudio para baixar (com progresso do download\n"
        "e da conversão em MP3).\n\n"
        "Comandos:\n"
        "/pastas — mostra as pastas configuradas e se estão acessíveis\n"
        "/id — mostra seu ID de usuário do Telegram"
    )
    if cfg.converter_perguntar and caminho_ffmpeg(cfg):
        ajuda += (
            "\n\nSe você mandar um áudio que não seja <b>MP3</b> ou um vídeo que não "
            "seja <b>MP4</b>, eu pergunto se quer converter (o Holyrics funciona melhor "
            "com esses formatos)."
        )
    if cfg.minio_configurado():
        ajuda += (
            "\n/enviar [categoria] — gera um link para enviar um arquivo grande "
            "direto ao servidor, sem passar pelo Telegram"
        )
    await update.effective_message.reply_text(ajuda, parse_mode=ParseMode.HTML)


async def cmd_pastas(update: Update, cfg: Config) -> None:
    linhas = ["<b>Pastas configuradas</b>"]
    for cat, pasta in cfg.pastas.items():
        icone, detalhe = estado_pasta(pasta)
        linhas.append(
            f"{icone} <b>{html.escape(cat)}</b> — {html.escape(detalhe)}\n"
            f"<code>{html.escape(str(pasta))}</code>"
        )
    if cfg.pdf_converter:
        linhas.append("ℹ️ PDF → PowerPoint ativado (pasta <b>apresentacao</b>)")
    if cfg.yt_ativo:
        ff = caminho_ffmpeg(cfg)
        linhas.append(
            f"ℹ️ YouTube ativado — vídeo até {cfg.yt_altura_maxima}p, "
            f"MP3 {cfg.yt_mp3_kbps} kbps, limite {cfg.yt_duracao_max_min} min\n"
            f"ffmpeg: {'✅ ' + html.escape(ff) if ff else '❌ não encontrado (sem MP3)'}"
        )
    await update.effective_message.reply_text("\n".join(linhas), parse_mode=ParseMode.HTML)


async def cmd_id(update: Update, cfg: Config) -> None:
    u = update.effective_user
    await update.effective_message.reply_text(
        f"Seu ID: <code>{u.id}</code>", parse_mode=ParseMode.HTML
    )


# ---------------------------------------------------------------------------
# Copiar do servidor local do Bot API (arquivos grandes, sem limite de 20 MB)
# ---------------------------------------------------------------------------

async def copiar_local(origem: Path, destino: Path, progresso: Progresso) -> int:
    """Copia um arquivo local (servidor local do Bot API) informando progresso."""
    feito = 0
    with open(origem, "rb") as entrada, open(destino, "wb") as saida:
        while True:
            bloco = await asyncio.to_thread(entrada.read, 1024 * 1024)
            if not bloco:
                break
            saida.write(bloco)
            feito += len(bloco)
            await progresso(feito)
    return feito


async def baixar_do_minio(store, chave: str, destino: Path, mensagem,
                          titulo_html: str, total: int, intervalo: float) -> Path:
    """Baixa do MinIO mostrando o progresso na mensagem (callback roda em thread)."""
    estado = {"feito": 0}

    def callback(bytes_transferidos: int) -> None:
        estado["feito"] += int(bytes_transferidos or 0)

    parar = asyncio.Event()

    def texto(feito: int) -> str:
        if total:
            pct = min(99, int(feito * 100 / total))
            return (f"⬇️ Recebendo <b>{titulo_html}</b> do servidor…\n"
                    f"{pct}%  ({fmt_bytes(feito)} de {fmt_bytes(total)})")
        return f"⬇️ Recebendo <b>{titulo_html}</b> do servidor… {fmt_bytes(feito)}"

    async def vigia():
        ultimo = None
        while not parar.is_set():
            atual = texto(estado["feito"])
            if atual != ultimo:
                ultimo = atual
                try:
                    await mensagem.edit_text(atual, parse_mode=ParseMode.HTML)
                except Exception as e:  # noqa: BLE001
                    LOG.debug("Falha ao atualizar o progresso do MinIO: %s", e)
            try:
                await asyncio.wait_for(parar.wait(), intervalo)
            except asyncio.TimeoutError:
                pass
        # uma última atualização com o total final
        try:
            await mensagem.edit_text(texto(estado["feito"]), parse_mode=ParseMode.HTML)
        except Exception:  # noqa: BLE001
            pass

    tarefa = asyncio.create_task(vigia())
    try:
        await asyncio.to_thread(store.baixar, chave, destino, callback)
    finally:
        parar.set()
        await tarefa
    return destino


# ---------------------------------------------------------------------------
# Baixar arquivos a partir de um link colado no chat
# ---------------------------------------------------------------------------

def achar_links_http(texto: str) -> list[str]:
    """Extrai links http(s) de um texto (sem repetir)."""
    vistos: set[str] = set()
    urls: list[str] = []
    for achado in re.finditer(r"https?://[^\s<>\"']+", texto or ""):
        url = achado.group(0).rstrip(").,;]}")
        if url not in vistos:
            vistos.add(url)
            urls.append(url)
    return urls


def host_bloqueado(url: str, cfg: Config) -> bool:
    if not cfg.links_bloqueados:
        return False
    host = (urlparse(url).hostname or "").lower()
    return any(host == b or host.endswith("." + b) for b in cfg.links_bloqueados)


def nome_do_link(url_final: str, headers) -> str:
    """Descobre um nome de arquivo razoável a partir do link/resposta."""
    disposicao = headers.get("content-disposition", "") or ""
    achado = re.search(r"filename\*?=(?:UTF-8''|\"?)([^\";]+)", disposicao, re.IGNORECASE)
    if achado:
        nome = unquote(achado.group(1).strip().strip('"'))
        if nome:
            return nome_seguro(nome)
    nome = unquote(Path(urlparse(url_final).path).name)
    if nome and "." in nome:
        return nome_seguro(nome)
    tipo = (headers.get("content-type", "") or "").split(";")[0].strip().lower()
    return nome_seguro((nome or "arquivo") + MIME_EXT.get(tipo, ""))


async def baixar_link(url: str, cfg: Config, mensagem) -> Path:
    """Baixa um link http(s) para a pasta da categoria e devolve o destino final."""
    timeout = httpx.Timeout(connect=30.0, read=120.0, write=60.0, pool=30.0)
    limite = cfg.links_max_mb * 1024 * 1024 if cfg.links_max_mb else 0
    temporario = Path(tempfile.mkdtemp(prefix="hl-link-"))
    try:
        async with httpx.AsyncClient(timeout=timeout, follow_redirects=True) as cliente:
            async with cliente.stream("GET", url) as resposta:
                resposta.raise_for_status()
                total = int(resposta.headers.get("content-length") or 0)
                if limite and total and total > limite:
                    raise RuntimeError(
                        f"o arquivo tem {fmt_bytes(total)} e passa do limite de "
                        f"{cfg.links_max_mb} MB para links."
                    )
                nome = nome_do_link(str(resposta.url), resposta.headers)
                parcial = temporario / nome
                progresso = Progresso(
                    mensagem, f"<b>{html.escape(nome)}</b>", total, cfg.progresso_intervalo
                )
                feito = 0
                with open(parcial, "wb") as saida:
                    async for bloco in resposta.aiter_bytes(512 * 1024):
                        saida.write(bloco)
                        feito += len(bloco)
                        if limite and feito > limite:
                            raise RuntimeError(
                                f"o arquivo passou do limite de {cfg.links_max_mb} MB para links."
                            )
                        await progresso(feito)

        if eh_pdf(nome) and cfg.pdf_converter:
            pasta = cfg.pastas.get("apresentacao") or cfg.pastas.get("outros")
            if pasta is None:
                raise RuntimeError("nenhuma pasta de apresentação configurada")
            destino = nome_livre(pasta / (Path(nome).stem + ".pptx"))
            destino.parent.mkdir(parents=True, exist_ok=True)
            await mensagem.edit_text(
                f"🔄 Convertendo <b>{html.escape(nome)}</b> para PowerPoint…",
                parse_mode=ParseMode.HTML,
            )
            try:
                await asyncio.to_thread(pdf_para_pptx, parcial, destino, cfg)
            except Exception as e:  # noqa: BLE001 — salva o PDF original
                LOG.exception("Falha ao converter PDF do link %s", url)
                destino = nome_livre(pasta / nome)
                shutil.move(str(parcial), str(destino))
                raise RuntimeError(
                    f"baixei o PDF, mas a conversão para PowerPoint falhou ({e}); "
                    f"salvei o original em {destino}"
                )
            if cfg.pdf_manter_original:
                shutil.move(str(parcial), str(nome_livre(pasta / nome)))
            return destino

        categoria = categoria_de(nome)
        pasta = cfg.pastas.get(categoria) or cfg.pastas.get("outros")
        if pasta is None:
            raise RuntimeError(f"nenhuma pasta configurada para '{categoria}'")
        destino = nome_livre(pasta / nome)
        destino.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(parcial), str(destino))
        return destino
    finally:
        shutil.rmtree(temporario, ignore_errors=True)


async def tratar_links_arquivo(update: Update, context: ContextTypes.DEFAULT_TYPE,
                               urls: list[str]) -> None:
    cfg: Config = context.bot_data["cfg"]
    msg = update.effective_message
    aceitos = [u for u in urls if not host_bloqueado(u, cfg)]
    if not aceitos:
        await msg.reply_text("🔒 Não baixo arquivos desse endereço.")
        return
    for url in aceitos[:3]:
        aviso = await msg.reply_text("🔎 Baixando do link…")
        try:
            destino = await baixar_link(url, cfg, aviso)
        except Exception as e:  # noqa: BLE001
            LOG.warning("Falha ao baixar o link %s: %s", url, e)
            await aviso.edit_text(
                f"❌ Não consegui baixar esse link:\n{html.escape(str(e))}",
                parse_mode=ParseMode.HTML,
            )
            continue
        await aviso.edit_text(
            f"✅ <b>{html.escape(destino.name)}</b>\n"
            f"📁 Salvo em:\n<code>{html.escape(str(destino))}</code>\n"
            f"📦 {fmt_bytes(destino.stat().st_size)}",
            parse_mode=ParseMode.HTML,
        )


# ---------------------------------------------------------------------------
# Arquivos grandes: upload pelo MinIO
# ---------------------------------------------------------------------------

def montar_link_minio(cfg: Config, dados: dict) -> str:
    if not cfg.minio_pagina_upload:
        return ""
    payload = {"url": dados["url"], "fields": dados.get("fields") or {}}
    codificado = base64.urlsafe_b64encode(
        json.dumps(payload).encode("utf-8")
    ).decode("ascii").rstrip("=")
    return f"{cfg.minio_pagina_upload}#{codificado}"


async def publicar_link_upload(cfg: Config, store, dados: dict) -> tuple[str, str | None]:
    """Publica os dados do POST num JSON do MinIO e devolve um link curto.

    O link fica `.../uploader.html#<token>` em vez de carregar a assinatura toda.
    Sem bucket público configurado, cai no link longo (dados no #).
    """
    if not cfg.minio_pagina_upload:
        return "", None
    if not cfg.minio_bucket_publico:
        return montar_link_minio(cfg, dados), None
    token = secrets.token_hex(8)
    chave_json = f"links/{token}.json"
    payload = {"url": dados["url"], "fields": dados.get("fields") or {}}
    await asyncio.to_thread(store.guardar_json, chave_json, payload, cfg.minio_bucket_publico)
    return f"{cfg.minio_pagina_upload}#{token}", chave_json


def salvar_para_pasta(origem: Path, nome: str, cfg: Config,
                      categoria_forcada: str | None = None) -> Path:
    """Move um arquivo baixado para a pasta da categoria (converte PDF se preciso)."""
    if eh_pdf(nome) and cfg.pdf_converter:
        pasta = cfg.pastas.get("apresentacao") or cfg.pastas.get("outros")
        if pasta is None:
            raise RuntimeError("nenhuma pasta de apresentação configurada")
        destino = nome_livre(pasta / (Path(nome).stem + ".pptx"))
        destino.parent.mkdir(parents=True, exist_ok=True)
        pdf_para_pptx(origem, destino, cfg)
        if cfg.pdf_manter_original:
            shutil.move(str(origem), str(nome_livre(pasta / nome)))
        return destino
    categoria = categoria_forcada or categoria_de(nome)
    pasta = cfg.pastas.get(categoria) or cfg.pastas.get("outros")
    if pasta is None:
        raise RuntimeError(f"nenhuma pasta configurada para '{categoria}'")
    destino = nome_livre(pasta / nome)
    destino.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(origem), str(destino))
    return destino


# ---------------------------------------------------------------------------
# Conversão de mídia para os formatos que o Holyrics prefere (mp3 / mp4)
# ---------------------------------------------------------------------------

def alvo_conversao(nome: str, cfg: Config) -> str | None:
    """Diz para qual formato perguntar a conversão (ou None se já está certo)."""
    if not cfg.converter_perguntar or not ffmpeg_disponivel(cfg):
        return None
    sufixo = Path(nome).suffix.lower()
    categoria = categoria_de(nome)
    if categoria == "audio" and sufixo != ".mp3":
        return "mp3"
    if categoria == "video" and sufixo != ".mp4":
        return "mp4"
    return None


def converter_midia(origem: Path, destino: Path, alvo: str, cfg: Config) -> Path:
    """Converte áudio para mp3 ou vídeo para mp4 usando o ffmpeg."""
    ffmpeg = caminho_ffmpeg(cfg)
    if not ffmpeg:
        raise RuntimeError("ffmpeg não encontrado para converter")
    if alvo == "mp3":
        comando = [ffmpeg, "-y", "-i", str(origem), "-vn",
                   "-codec:a", "libmp3lame", "-b:a", f"{cfg.converter_audio_kbps}k", str(destino)]
    else:
        comando = [ffmpeg, "-y", "-i", str(origem), "-c:v", "libx264", "-preset", "veryfast",
                   "-crf", "20", "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart",
                   str(destino)]
    resultado = subprocess.run(comando, capture_output=True, text=True)
    if resultado.returncode != 0:
        raise RuntimeError(f"ffmpeg falhou: {resultado.stderr.strip()[-300:]}")
    return destino


async def oferecer_conversao(app: Application, chat_id: int, mensagem, origem: Path,
                             nome: str, alvo: str, usuario: int) -> None:
    """Guarda o arquivo e pergunta se o usuário quer converter (mp3/mp4)."""
    token = secrets.token_hex(5)
    staging = Path(tempfile.mkdtemp(prefix="hl-conv-"))
    destino_staging = staging / nome_seguro(nome)
    shutil.move(str(origem), str(destino_staging))

    app.bot_data.setdefault("conversoes", {})[token] = {
        "caminho": destino_staging,
        "nome": nome,
        "alvo": alvo,
        "user": usuario,
        "chat": chat_id,
        "staging": staging,
        "mensagem": mensagem,
    }
    rotulo = "MP3" if alvo == "mp3" else "MP4"
    tipo = "áudio" if alvo == "mp3" else "vídeo"
    teclado = InlineKeyboardMarkup([[
        InlineKeyboardButton(f"✅ Converter para {rotulo}", callback_data=f"conv|{token}|1"),
        InlineKeyboardButton("Salvar como está", callback_data=f"conv|{token}|0"),
    ]])
    await editar(
        mensagem,
        f"📥 <b>{html.escape(nome)}</b> ({fmt_bytes(destino_staging.stat().st_size)}) recebido.\n"
        f"O Holyrics funciona melhor com <b>{rotulo}</b> para {tipo}. Quer converter?",
        reply_markup=teclado,
    )
    asyncio.create_task(_expirar_conversao(app, token))


async def _expirar_conversao(app: Application, token: str) -> None:
    """Se o usuário não responder, salva o original (não perde o arquivo)."""
    await asyncio.sleep(3600)
    pendentes: dict = app.bot_data.setdefault("conversoes", {})
    pendente = pendentes.pop(token, None)
    if pendente is None:
        return
    cfg: Config = app.bot_data["cfg"]
    try:
        destino = await asyncio.to_thread(
            salvar_para_pasta, pendente["caminho"], pendente["nome"], cfg
        )
        await editar(
            pendente["mensagem"],
            f"⌛ Passou do tempo, mas salvei <b>{html.escape(pendente['nome'])}</b> "
            f"como estava:\n<code>{html.escape(str(destino))}</code>",
        )
    except Exception as e:  # noqa: BLE001
        LOG.warning("Falha ao salvar conversão expirada: %s", e)
    finally:
        shutil.rmtree(pendente["staging"], ignore_errors=True)


async def botao_converter(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    cfg: Config = context.bot_data["cfg"]
    consulta = update.callback_query
    partes = (consulta.data or "").split("|")
    pendentes: dict = context.bot_data.setdefault("conversoes", {})
    pendente = pendentes.pop(partes[1], None) if len(partes) == 3 else None
    mensagem = consulta.message

    if pendente is None:
        await consulta.answer("Essa opção expirou.", show_alert=True)
        await editar(mensagem, "⌛ Opção expirada.")
        return
    if consulta.from_user.id != pendente["user"]:
        await consulta.answer("Só quem enviou o arquivo pode escolher.", show_alert=True)
        return

    await consulta.answer()
    try:
        await mensagem.edit_reply_markup(None)
    except Exception:  # noqa: BLE001
        pass

    converter = partes[2] == "1"
    origem = Path(pendente["caminho"])
    nome = pendente["nome"]
    try:
        if converter:
            alvo = pendente["alvo"]
            novo_nome = nome_seguro(Path(nome).with_suffix("." + alvo).name)
            saida = origem.with_name(novo_nome)
            await editar(mensagem, f"🔄 Convertendo <b>{html.escape(nome)}</b> para {alvo.upper()}…")
            await asyncio.to_thread(converter_midia, origem, saida, alvo, cfg)
            origem, nome = saida, novo_nome
        destino = await asyncio.to_thread(salvar_para_pasta, origem, nome, cfg)
        await editar(
            mensagem,
            f"✅ <b>{html.escape(nome)}</b>\n"
            f"📁 Salvo em:\n<code>{html.escape(str(destino))}</code>\n"
            f"📦 {fmt_bytes(destino.stat().st_size)}",
        )
    except Exception as e:  # noqa: BLE001
        LOG.exception("Falha ao converter/salvar %s", nome)
        await editar(mensagem, f"❌ Falha com <b>{html.escape(nome)}</b>:\n{html.escape(str(e))}")
    finally:
        shutil.rmtree(pendente["staging"], ignore_errors=True)


async def oferecer_upload(update: Update, context: ContextTypes.DEFAULT_TYPE,
                          msg, nome: str, tamanho: int) -> None:
    """Oferece um link do MinIO para o usuário enviar um arquivo grande."""
    cfg: Config = context.bot_data["cfg"]
    store = context.bot_data.get("minio")
    if store is None or not cfg.minio_configurado():
        await msg.reply_text(
            f"❌ <b>{html.escape(nome)}</b> tem {fmt_bytes(tamanho)} e passa do limite de "
            f"{cfg.max_mb} MB que o Telegram entrega para bots.\n"
            "Para receber arquivos maiores, ative o MinIO no config.json.",
            parse_mode=ParseMode.HTML,
        )
        return

    token = secrets.token_hex(6)
    chave = f"{cfg.minio_prefixo}{token}-{nome_seguro(nome)}"
    redirect = None
    if cfg.minio_pagina_upload:
        separador = "&" if "?" in cfg.minio_pagina_upload else "?"
        redirect = f"{cfg.minio_pagina_upload}{separador}ok=1"
    dados = await asyncio.to_thread(store.dados_upload, chave, redirect)
    link, link_json = await publicar_link_upload(cfg, store, dados)
    if not link:
        link = await asyncio.to_thread(store.url_put, chave)

    context.bot_data.setdefault("upload_pendentes", {})[token] = {
        "chave": chave,
        "nome": nome,
        "user": update.effective_user.id,
        "chat": msg.chat_id,
        "link_json": link_json,
    }

    if cfg.minio_pagina_upload:
        instrucao = f"🔗 <b>Abra o link e envie o arquivo:</b>\n{html.escape(link)}"
    else:
        instrucao = (
            "🔗 <b>Envie o arquivo</b> (num computador) com:\n"
            f"<code>curl -T \"caminho/do/arquivo\" \"{html.escape(link)}\"</code>"
        )
    await msg.reply_text(
        f"📦 <b>{html.escape(nome)}</b> tem {fmt_bytes(tamanho)} — acima do limite de "
        f"{cfg.max_mb} MB do Telegram.\n\n{instrucao}\n\n"
        f"Assim que o envio terminar, eu salvo na pasta certa automaticamente "
        f"(aguardo até {cfg.minio_espera_min} min).",
        parse_mode=ParseMode.HTML,
        disable_web_page_preview=True,
    )
    asyncio.create_task(esperar_upload(context.application, token, msg.chat_id))


async def esperar_upload(app: Application, token: str, chat_id: int) -> None:
    """Espera o objeto aparecer no MinIO, baixa, salva e apaga."""
    cfg: Config = app.bot_data["cfg"]
    store = app.bot_data.get("minio")
    pendentes: dict = app.bot_data.setdefault("upload_pendentes", {})
    pendente = pendentes.get(token)
    if store is None or pendente is None:
        return

    limite = time.monotonic() + cfg.minio_espera_min * 60
    tamanho = 0
    while time.monotonic() < limite:
        try:
            tamanho = await asyncio.to_thread(store.tamanho, pendente["chave"])
        except Exception as e:  # noqa: BLE001
            LOG.warning("Falha ao consultar o MinIO: %s", e)
            tamanho = 0
        if tamanho > 0:
            break
        await asyncio.sleep(max(3, cfg.minio_intervalo_s))

    pendentes.pop(token, None)
    if pendente.get("link_json"):
        await asyncio.to_thread(store.apagar, pendente["link_json"], cfg.minio_bucket_publico)
    if tamanho <= 0:
        await app.bot.send_message(
            chat_id,
            f"⌛ Não recebi <b>{html.escape(pendente['nome'])}</b> a tempo. "
            "Mande o arquivo de novo para gerar um link novo.",
            parse_mode=ParseMode.HTML,
        )
        return

    nome = pendente["nome"]
    estado = await app.bot.send_message(
        chat_id,
        f"⬇️ Recebendo <b>{html.escape(nome)}</b> do servidor… 0%",
        parse_mode=ParseMode.HTML,
    )
    temporario = Path(tempfile.mkdtemp(prefix="hl-minio-"))
    try:
        baixado = temporario / nome_seguro(nome)
        await baixar_do_minio(store, pendente["chave"], baixado, estado,
                              html.escape(nome), tamanho, cfg.progresso_intervalo)
        await asyncio.to_thread(store.apagar, pendente["chave"])
        alvo = alvo_conversao(nome, cfg)
        if alvo:
            await oferecer_conversao(app, chat_id, estado, baixado, nome, alvo, pendente["user"])
        else:
            destino = await asyncio.to_thread(salvar_para_pasta, baixado, nome, cfg)
            await estado.edit_text(
                f"✅ <b>{html.escape(nome)}</b>\n"
                f"📁 Salvo em:\n<code>{html.escape(str(destino))}</code>\n"
                f"📦 {fmt_bytes(destino.stat().st_size)}",
                parse_mode=ParseMode.HTML,
            )
    except Exception as e:  # noqa: BLE001
        LOG.exception("Falha ao salvar o arquivo vindo do MinIO")
        await estado.edit_text(
            f"❌ Recebi <b>{html.escape(nome)}</b>, mas falhei ao salvar: "
            f"{html.escape(str(e))}",
            parse_mode=ParseMode.HTML,
        )
    finally:
        shutil.rmtree(temporario, ignore_errors=True)


async def esperar_prefixo(app: Application, token: str, chat_id: int) -> None:
    """Espera qualquer arquivo aparecer numa pasta do MinIO (comando /enviar)."""
    cfg: Config = app.bot_data["cfg"]
    store = app.bot_data.get("minio")
    pendentes: dict = app.bot_data.setdefault("upload_pendentes", {})
    pendente = pendentes.get(token)
    if store is None or pendente is None:
        return

    prefixo = pendente["prefixo"]
    limite = time.monotonic() + cfg.minio_espera_min * 60
    chave = None
    while time.monotonic() < limite:
        try:
            achados = [c for c in await asyncio.to_thread(store.listar, prefixo)
                       if not c.endswith("/")]
        except Exception as e:  # noqa: BLE001
            LOG.warning("Falha ao listar o MinIO: %s", e)
            achados = []
        if achados:
            chave = achados[0]
            break
        await asyncio.sleep(max(3, cfg.minio_intervalo_s))

    pendentes.pop(token, None)
    if pendente.get("link_json"):
        await asyncio.to_thread(store.apagar, pendente["link_json"], cfg.minio_bucket_publico)
    if not chave:
        await app.bot.send_message(
            chat_id,
            "⌛ Não recebi nenhum arquivo a tempo. Use /enviar de novo para pegar um link novo.",
            parse_mode=ParseMode.HTML,
        )
        return

    nome = chave.rsplit("/", 1)[-1]
    try:
        tamanho = await asyncio.to_thread(store.tamanho, chave)
    except Exception:  # noqa: BLE001
        tamanho = 0
    estado = await app.bot.send_message(
        chat_id,
        f"⬇️ Recebendo <b>{html.escape(nome)}</b> do servidor… 0%",
        parse_mode=ParseMode.HTML,
    )
    temporario = Path(tempfile.mkdtemp(prefix="hl-minio-"))
    try:
        baixado = temporario / nome_seguro(nome)
        await baixar_do_minio(store, chave, baixado, estado,
                              html.escape(nome), tamanho, cfg.progresso_intervalo)
        await asyncio.to_thread(store.apagar, chave)
        alvo = alvo_conversao(nome, cfg)
        if alvo:
            await oferecer_conversao(app, chat_id, estado, baixado, nome, alvo, pendente["user"])
        else:
            destino = await asyncio.to_thread(
                salvar_para_pasta, baixado, nome_seguro(nome), cfg, pendente.get("categoria")
            )
            await estado.edit_text(
                f"✅ <b>{html.escape(nome)}</b>\n"
                f"📁 Salvo em:\n<code>{html.escape(str(destino))}</code>\n"
                f"📦 {fmt_bytes(destino.stat().st_size)}",
                parse_mode=ParseMode.HTML,
            )
    except Exception as e:  # noqa: BLE001
        LOG.exception("Falha ao salvar o arquivo do /enviar")
        await estado.edit_text(
            f"❌ Recebi <b>{html.escape(nome)}</b>, mas falhei ao salvar: "
            f"{html.escape(str(e))}",
            parse_mode=ParseMode.HTML,
        )
    finally:
        shutil.rmtree(temporario, ignore_errors=True)


async def cmd_enviar(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Gera um link do MinIO para o usuário enviar um arquivo grande sem passar pelo Telegram."""
    cfg: Config = context.bot_data["cfg"]
    store = context.bot_data.get("minio")
    msg = update.effective_message
    if not autorizado(update, cfg):
        await msg.reply_text("🔒 Não autorizado.")
        return
    if store is None or not cfg.minio_configurado():
        await msg.reply_text("O MinIO não está ativado no config.json.")
        return
    if not cfg.minio_pagina_upload:
        await msg.reply_text(
            "Para usar /enviar, configure <code>minio.url_pagina_upload</code> no config.json "
            "(veja o README). Sem a página, mande o arquivo pelo chat e eu gero o link.",
            parse_mode=ParseMode.HTML,
        )
        return

    categoria = None
    if context.args:
        escolha = context.args[0].strip().lower()
        if escolha not in cfg.pastas:
            await msg.reply_text(
                "Categoria inválida. Use: " + ", ".join(sorted(cfg.pastas))
                + "\nEx.: <code>/enviar video</code>",
                parse_mode=ParseMode.HTML,
            )
            return
        categoria = escolha

    token = secrets.token_hex(6)
    prefixo = f"{cfg.minio_prefixo}{token}/"
    separador = "&" if "?" in cfg.minio_pagina_upload else "?"
    redirect = f"{cfg.minio_pagina_upload}{separador}ok=1"
    # ${filename} faz o MinIO usar o nome real do arquivo escolhido no navegador.
    dados = await asyncio.to_thread(store.dados_upload, prefixo + "${filename}", redirect)
    link, link_json = await publicar_link_upload(cfg, store, dados)

    context.bot_data.setdefault("upload_pendentes", {})[token] = {
        "prefixo": prefixo,
        "categoria": categoria,
        "user": update.effective_user.id,
        "chat": msg.chat_id,
        "nome": None,
        "link_json": link_json,
    }
    destino = f"<b>{categoria}</b>" if categoria else "da categoria do arquivo"
    await msg.reply_text(
        "📤 <b>Envie o arquivo grande por este link</b> — não precisa mandar no Telegram:\n"
        f"{html.escape(link)}\n\n"
        f"Assim que o envio terminar, eu salvo na pasta {destino} "
        f"(aguardo até {cfg.minio_espera_min} min).",
        parse_mode=ParseMode.HTML,
        disable_web_page_preview=True,
    )
    asyncio.create_task(esperar_prefixo(context.application, token, msg.chat_id))


async def receber(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    cfg: Config = context.bot_data["cfg"]
    msg = update.effective_message
    if msg is None:
        return

    if not autorizado(update, cfg):
        await msg.reply_text("🔒 Não autorizado.")
        LOG.warning("Mensagem de usuário não autorizado: %s", update.effective_user)
        return
    if msg.chat.type != "private" and not cfg.aceitar_grupos:
        return

    dados = extrair_arquivo(msg)
    if not dados:
        if msg.text and not msg.text.startswith("/"):
            links = achar_videos_youtube(msg.text)
            if links:
                await tratar_links(update, context, links)
                return
            if cfg.links_ativo:
                urls = achar_links_http(msg.text)
                if urls:
                    await tratar_links_arquivo(update, context, urls)
                    return
            await msg.reply_text(
                "Me envie um <b>arquivo</b> (áudio, vídeo, imagem, PDF ou PowerPoint), um "
                "<b>link do YouTube</b> ou um <b>link direto</b> de arquivo que eu salvo na "
                "pasta do Holyrics.",
                parse_mode=ParseMode.HTML,
            )
        return

    file_id, nome_bruto, tamanho = dados
    nome = nome_seguro(nome_bruto)
    titulo_html = f"<b>{html.escape(nome)}</b>"
    LOG.info("Recebido: %s (%s) de %s", nome, fmt_bytes(tamanho), update.effective_user.id)

    if cfg.max_mb and tamanho and tamanho > cfg.max_mb * 1024 * 1024:
        await oferecer_upload(update, context, msg, nome, tamanho)
        return

    # 1) mensagem única de status (será editada até o resultado final)
    estado = await msg.reply_text(
        template_estado(titulo_html, tamanho), parse_mode=ParseMode.HTML
    )
    try:
        await context.bot.send_chat_action(msg.chat_id, ChatAction.TYPING)
    except Exception:
        pass

    temporario = Path(tempfile.mkdtemp(prefix="hl-bot-"))
    try:
        try:
            telegram_file = await context.bot.get_file(file_id)
        except Exception as e:  # noqa: BLE001
            # O Telegram recusa getFile acima de 20 MB na API na nuvem.
            if "too big" in str(e).lower():
                try:
                    await estado.delete()
                except Exception:
                    pass
                await oferecer_upload(update, context, msg, nome, tamanho)
                return
            raise

        total = telegram_file.file_size or tamanho
        # Nas versões atuais do python-telegram-bot, get_file() já devolve
        # file_path como URL completa (ex.: https://api.telegram.org/file/bot<token>/...).
        # Em versões antigas (e nos testes) vem só o caminho relativo, então montamos
        # a URL apenas quando necessário — nunca prefixar duas vezes.
        # Com o servidor local do Bot API, file_path é um caminho de arquivo no disco.
        caminho_arquivo = telegram_file.file_path
        if not caminho_arquivo:
            raise RuntimeError("o Telegram não informou o caminho do arquivo (file_path vazio)")
        baixado_path = temporario / nome
        progresso = Progresso(estado, titulo_html, total, cfg.progresso_intervalo)
        origem_local = Path(caminho_arquivo)
        if origem_local.is_absolute() and origem_local.is_file():
            await copiar_local(origem_local, baixado_path, progresso)
        else:
            if caminho_arquivo.startswith(("http://", "https://")):
                url = caminho_arquivo
            else:
                url = f"{BASE_API}/file/bot{cfg.token}/{caminho_arquivo}"
            await baixar(url, baixado_path, progresso)

        # 1.5) áudio fora de MP3 / vídeo fora de MP4 -> pergunta se quer converter
        alvo = alvo_conversao(nome, cfg)
        if alvo:
            await oferecer_conversao(context.application, msg.chat_id, estado,
                                     baixado_path, nome, alvo, update.effective_user.id)
            return

        # 2) decide a pasta de destino (PDF converte antes de escolher destino)
        if eh_pdf(nome) and cfg.pdf_converter:
            pasta = cfg.pastas.get("apresentacao") or cfg.pastas.get("outros")
            if pasta is None:
                raise RuntimeError("nenhuma pasta de apresentação configurada")
            destino = nome_livre(pasta / (Path(nome).stem + ".pptx"))
            await estado.edit_text(
                f"🔄 <b>{html.escape(nome)}</b> baixado ({fmt_bytes(total)})\n"
                "Convertendo para PowerPoint…",
                parse_mode=ParseMode.HTML,
            )
            try:
                destino.parent.mkdir(parents=True, exist_ok=True)
                paginas = await asyncio.to_thread(pdf_para_pptx, baixado_path, destino, cfg)
            except Exception as e:
                LOG.exception("Falha ao converter PDF %s", nome)
                destino = nome_livre(pasta / nome)
                destino.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(str(baixado_path), str(destino))
                await estado.edit_text(
                    f"⚠️ <b>{html.escape(nome)}</b> baixado, mas a conversão para PowerPoint "
                    f"falhou ({html.escape(str(e))}).\nSalvei o PDF original em:\n"
                    f"<code>{html.escape(str(destino))}</code>",
                    parse_mode=ParseMode.HTML,
                )
                return
            if cfg.pdf_manter_original:
                shutil.move(str(baixado_path), str(nome_livre(pasta / nome)))
            texto = (
                f"✅ <b>{html.escape(nome)}</b>\n"
                f"🔄 Convertido para PowerPoint ({paginas} "
                f"{'slide' if paginas == 1 else 'slides'})\n"
                f"📁 Salvo em:\n<code>{html.escape(str(destino))}</code>\n"
                f"📦 {fmt_bytes(destino.stat().st_size)}"
            )
        else:
            pasta = cfg.pastas.get(categoria_de(nome)) or cfg.pastas.get("outros")
            if pasta is None:
                raise RuntimeError(f"nenhuma pasta configurada para '{categoria_de(nome)}'")
            destino = nome_livre(pasta / nome)
            destino.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(baixado_path), str(destino))
            texto = (
                f"✅ <b>{html.escape(nome)}</b>\n"
                f"📁 Salvo em:\n<code>{html.escape(str(destino))}</code>\n"
                f"📦 {fmt_bytes(destino.stat().st_size)}"
            )

        LOG.info("Salvo: %s", destino)
        await estado.edit_text(texto, parse_mode=ParseMode.HTML)

    except Exception as e:
        LOG.exception("Erro ao processar %s", nome)
        try:
            await estado.edit_text(
                f"❌ Falha ao receber <b>{html.escape(nome)}</b>:\n{html.escape(str(e))}",
                parse_mode=ParseMode.HTML,
            )
        except Exception:
            pass
    finally:
        shutil.rmtree(temporario, ignore_errors=True)


def configurar_log(cfg: Config) -> None:
    nivel = getattr(logging, cfg.log_nivel, logging.INFO)
    formato = logging.Formatter("%(asctime)s %(levelname)s %(name)s: %(message)s")

    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(formato)

    destino = Path(cfg.log_arquivo)
    if not destino.is_absolute():
        destino = RAIZ / destino
    arquivo = logging.handlers.RotatingFileHandler(
        destino, maxBytes=2 * 1024 * 1024, backupCount=3, encoding="utf-8"
    )
    arquivo.setFormatter(formato)

    logging.basicConfig(level=nivel, handlers=[console, arquivo])


def verificar(cfg: Config) -> None:
    """Mostra as pastas resolvidas (útil para conferir a instalação)."""
    print(f"Config: {cfg.caminho}")
    for cat, pasta in cfg.pastas.items():
        icone, detalhe = estado_pasta(pasta)
        print(f"  [{icone}] {cat:13} -> {pasta}   ({detalhe})")
    print(f"PDF -> PowerPoint: {'ativado' if cfg.pdf_converter else 'desativado'}")
    if cfg.yt_ativo:
        print(f"YouTube: ativado (vídeo até {cfg.yt_altura_maxima}p, MP3 {cfg.yt_mp3_kbps} kbps, "
              f"limite {cfg.yt_duracao_max_min} min)")
        ff = caminho_ffmpeg(cfg)
        print(f"  ffmpeg: {ff}" if ff else "  ffmpeg: NÃO ENCONTRADO (sem MP3 e sem juntar vídeo+áudio)")
        if not ff:
            print("          rode:  python baixar_ffmpeg.py")
        if cfg.yt_cookies:
            print(f"  cookies: {cfg.yt_cookies} ({'existe' if Path(cfg.yt_cookies).is_file() else 'ARQUIVO NÃO ENCONTRADO'})")
        try:
            import yt_dlp

            print(f"  yt-dlp: {yt_dlp.version.__version__}")
        except ImportError:
            print("  yt-dlp: NÃO INSTALADO")
    else:
        print("YouTube: desativado")
    print(f"Usuários autorizados: {sorted(cfg.usuarios) or 'TODOS (config sem lista)'}")
    if cfg.local_mode:
        print(f"Servidor local do Bot API: {cfg.base_url} (sem limite de 20 MB)")
    else:
        print(f"Telegram (nuvem): limite de {cfg.max_mb or '∞'} MB por arquivo recebido")
    print(f"Links diretos: {'ativados' if cfg.links_ativo else 'desativados'}"
          + (f" (limite {cfg.links_max_mb} MB)" if cfg.links_max_mb else " (sem limite)"))
    if cfg.minio_configurado():
        print(f"MinIO: ativado em {cfg.minio_endpoint} (bucket '{cfg.minio_bucket}')")
    else:
        print("MinIO: desativado (arquivos acima do limite não terão link de upload)")
    if cfg.pdf_converter:
        amostra = cfg.pastas.get("apresentacao")
        if amostra is not None:
            try:
                amostra.mkdir(parents=True, exist_ok=True)
                teste = amostra / ".escrita-teste"
                teste.write_text("ok", encoding="utf-8")
                teste.unlink()
                print("Escrita na pasta de apresentações: OK")
            except Exception as e:
                print(f"Escrita na pasta de apresentações: FALHOU ({e})")


async def publicar_comandos(app: Application) -> None:
    """Registra o menu de comandos que aparece ao digitar '/' no Telegram."""
    cfg: Config = app.bot_data["cfg"]
    comandos = [
        BotCommand("start", "Ajuda"),
        BotCommand("pastas", "Mostra as pastas configuradas"),
        BotCommand("id", "Mostra seu ID do Telegram"),
    ]
    if cfg.minio_configurado():
        comandos.append(BotCommand("enviar", "Link para enviar um arquivo grande"))
    try:
        await app.bot.set_my_commands(comandos)
        LOG.info("Menu de comandos publicado (%d comandos).", len(comandos))
    except Exception as e:  # noqa: BLE001
        LOG.warning("Não consegui publicar o menu de comandos: %s", e)


def main() -> None:
    parser = argparse.ArgumentParser(description="Bot do Telegram x Holyrics")
    parser.add_argument("--config", default="config.json", help="caminho do arquivo de configuração")
    parser.add_argument("--verificar", action="store_true", help="confere a configuração e sai")
    args = parser.parse_args()

    caminho = Path(args.config)
    if not caminho.is_absolute():
        caminho = RAIZ / caminho
    cfg = Config(caminho)
    configurar_log(cfg)

    if args.verificar:
        verificar(cfg)
        return

    logging.getLogger("httpx").setLevel(logging.WARNING)
    LOG.info("Iniciando bot (pastas: %s)", {k: str(v) for k, v in cfg.pastas.items()})

    # Raiz usada para montar a URL de download dos arquivos. Com o servidor local
    # do Bot API, aponta para ele; na nuvem, para api.telegram.org.
    global BASE_API
    if cfg.base_file_url.endswith("/file/bot"):
        BASE_API = cfg.base_file_url[: -len("/file/bot")]
    else:
        BASE_API = cfg.base_file_url
    LOG.info("API do Telegram: %s (arquivos: %s, modo local: %s)",
             cfg.base_url, BASE_API, cfg.local_mode)

    store = None
    if cfg.minio_configurado():
        try:
            store = MinioStore(cfg)
            store.garantir_bucket()
            if cfg.minio_bucket_publico:
                antigos = store.limpar_antigos("links/", cfg.minio_expira_min, cfg.minio_bucket_publico)
                if antigos:
                    LOG.info("Links de upload antigos removidos: %d", antigos)
            LOG.info("MinIO pronto: %s (bucket '%s')", store.endpoint, store.bucket)
        except Exception:
            LOG.exception("MinIO configurado, mas inacessível: "
                          "upload de arquivos grandes ficará indisponível.")
            store = None

    construtor = (
        Application.builder()
        .token(cfg.token)
        .concurrent_updates(True)
        .read_timeout(60)
        .write_timeout(300)
        .base_url(cfg.base_url)
        .base_file_url(cfg.base_file_url)
    )
    if cfg.local_mode:
        construtor = construtor.local_mode(True)
    app = construtor.build()
    app.bot_data["cfg"] = cfg
    app.bot_data["minio"] = store
    app.bot_data["upload_pendentes"] = {}
    app.post_init = publicar_comandos

    app.add_handler(CommandHandler("start", lambda u, c: cmd_start(u, cfg)))
    app.add_handler(CommandHandler("help", lambda u, c: cmd_start(u, cfg)))
    app.add_handler(CommandHandler("pastas", lambda u, c: cmd_pastas(u, cfg)))
    app.add_handler(CommandHandler("id", lambda u, c: cmd_id(u, cfg)))
    app.add_handler(CommandHandler("enviar", cmd_enviar))
    app.add_handler(CommandHandler("upload", cmd_enviar))
    app.add_handler(CallbackQueryHandler(botao_youtube, pattern=r"^yt\|"))
    app.add_handler(CallbackQueryHandler(botao_converter, pattern=r"^conv\|"))
    app.add_handler(MessageHandler(filters.ALL & ~filters.COMMAND, receber))

    LOG.info("Bot no ar. Ctrl+C para encerrar.")
    app.run_polling(drop_pending_updates=True, allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
